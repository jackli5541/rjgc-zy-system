from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from datetime import datetime, timedelta, timezone

UTC = timezone.utc
from decimal import Decimal
from uuid import UUID
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

from app.database import SessionLocal
from app.main import app, parse_roster
from app.grading import final_score
from app.models import Assignment, AuditLog, Grade, PeerReview, ReviewAssignment, ReviewCampaign, Submission, SubmissionDocument, SubmissionVersion, Team, TeamRequest
from app.worker import process_auto_review, process_due_campaign


def login(account: str, password: str, role: str):
    client = TestClient(app)
    response = client.post("/api/v1/auth/login", json={"account": account, "password": password, "role": role})
    assert response.status_code == 200, response.text
    return client, {"X-CSRF-Token": response.json()["csrf_token"]}


def test_published_assignment_can_change_class():
    teacher, headers = login("teacher", "123456", "teacher")
    first = teacher.post("/api/v1/classes", headers=headers, json={"semester": "班级调整", "name": "原教学班"}).json()
    second = teacher.post("/api/v1/classes", headers=headers, json={"semester": "班级调整", "name": "新教学班"}).json()
    assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": first["id"], "title": "已发布作业", "description": "调整教学班", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    url = f"/api/v1/assignments/{assignment['id']}"
    missing = teacher.patch(url, headers=headers, json={"class_id": "00000000-0000-0000-0000-000000000000", "version": assignment["version"]})
    assert missing.status_code == 404
    changed = teacher.patch(url, headers=headers, json={"class_id": second["id"], "version": assignment["version"]})
    assert changed.status_code == 200, changed.text
    assert changed.json()["class_id"] == second["id"]
    assert changed.json()["status"] == "PUBLISHED"
    assert teacher.get(f"/api/v1/assignments?class_id={first['id']}").json()["total"] == 0
    assert teacher.get(f"/api/v1/assignments?class_id={second['id']}").json()["total"] == 1
    stale = teacher.patch(url, headers=headers, json={"class_id": first["id"], "version": assignment["version"]})
    assert stale.status_code == 409
    teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"status": "ARCHIVED", "version": first["version"]})
    archived = teacher.patch(url, headers=headers, json={"class_id": first["id"], "version": changed.json()["version"]})
    assert archived.status_code == 409
    assert teacher.get(f"/api/v1/assignments?class_id={second['id']}").json()["items"][0]["version"] == changed.json()["version"]


def test_assignment_class_change_preserves_existing_submissions():
    teacher, headers = login("teacher", "123456", "teacher")
    first = teacher.post("/api/v1/classes", headers=headers, json={"semester": "班级保护", "name": "原教学班"}).json()
    second = teacher.post("/api/v1/classes", headers=headers, json={"semester": "班级保护", "name": "新教学班"}).json()
    member = teacher.post(f"/api/v1/classes/{first['id']}/members", headers=headers, json={"student_no": "20998881", "name": "提交学生"}).json()
    assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": first["id"], "title": "已有提交作业", "description": "保护提交记录", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    with SessionLocal() as db:
        db.add(Submission(assignment_id=UUID(assignment["id"]), owner_user_id=UUID(member["id"]), status="DRAFT"))
        db.commit()
    url = f"/api/v1/assignments/{assignment['id']}"
    locked = teacher.patch(url, headers=headers, json={"class_id": second["id"], "version": assignment["version"]})
    assert locked.status_code == 409, locked.text
    unchanged = teacher.patch(url, headers=headers, json={"class_id": first["id"], "title": "原班正常编辑", "version": assignment["version"]})
    assert unchanged.status_code == 200, unchanged.text
    assert unchanged.json()["class_id"] == first["id"]


def test_team_leader_can_close_recruitment():
    teacher, headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=headers, json={"semester": "招募测试", "name": "招募教学班"}).json()
    students = []
    for index in range(3):
        account = f"2099777{index}"
        teacher.post(f"/api/v1/classes/{course['id']}/members", headers=headers, json={"student_no": account, "name": f"招募学生{index}"})
        students.append(login(account, account, "student"))
    leader, leader_headers = students[0]
    applicant, applicant_headers = students[1]
    other, other_headers = students[2]
    team = leader.post("/api/v1/teams", headers=leader_headers, json={"class_id": course["id"], "name": "招募小组"}).json()
    url = f"/api/v1/teams/{team['id']}/close-recruitment"
    application = applicant.post(f"/api/v1/teams/{team['id']}/applications", headers=applicant_headers).json()
    other_id = other.get("/api/v1/auth/session").json()["user"]["id"]
    invitation = leader.post(f"/api/v1/teams/{team['id']}/invitations", headers=leader_headers, json={"student_id": other_id})
    assert invitation.status_code == 201, invitation.text
    assert other.post(url, headers=other_headers).status_code == 403
    assert teacher.post(url, headers=headers).status_code == 403
    assert leader.post(url).status_code == 403
    closed = leader.post(url, headers=leader_headers)
    assert closed.status_code == 200, closed.text
    assert closed.json()["open_recruitment"] is False
    assert closed.json()["version"] == team["version"] + 1
    cancelled = leader.get(f"/api/v1/team-requests?class_id={course['id']}").json()["items"]
    assert len(cancelled) == 2 and all(item["status"] == "CANCELLED" for item in cancelled)
    assert closed.json()["pending_count"] == 0
    repeated = leader.post(url, headers=leader_headers)
    assert repeated.json()["version"] == closed.json()["version"]
    blocked = other.post(f"/api/v1/teams/{team['id']}/applications", headers=other_headers)
    assert blocked.status_code == 409
    assert blocked.json()["code"] == "TEAM_NOT_OPEN"
    invite_blocked = leader.post(f"/api/v1/teams/{team['id']}/invitations", headers=leader_headers, json={"student_id": other_id})
    assert invite_blocked.status_code == 409 and invite_blocked.json()["code"] == "TEAM_NOT_OPEN"
    accept_blocked = other.post(f"/api/v1/team-requests/{invitation.json()['id']}/respond?decision=APPROVED", headers=other_headers)
    assert accept_blocked.status_code == 409 and accept_blocked.json()["code"] == "INVITATION_NOT_PENDING"
    approval_blocked = leader.post(f"/api/v1/team-requests/{application['id']}/decision?decision=APPROVED", headers=leader_headers)
    assert approval_blocked.status_code == 409 and approval_blocked.json()["code"] == "REQUEST_NOT_PENDING"
    assert leader.get(f"/api/v1/teams/{team['id']}").json()["member_count"] == 1
    open_url = f"/api/v1/teams/{team['id']}/open-recruitment"
    assert other.post(open_url, headers=other_headers).status_code == 403
    assert teacher.post(open_url, headers=headers).status_code == 403
    assert leader.post(open_url).status_code == 403
    reopened = leader.post(open_url, headers=leader_headers)
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["open_recruitment"] is True
    assert reopened.json()["version"] == closed.json()["version"] + 1
    assert leader.post(open_url, headers=leader_headers).json()["version"] == reopened.json()["version"]
    assert leader.post(f"/api/v1/team-requests/{application['id']}/decision?decision=APPROVED", headers=leader_headers).status_code == 409
    assert other.post(f"/api/v1/team-requests/{invitation.json()['id']}/respond?decision=APPROVED", headers=other_headers).status_code == 409
    application = applicant.post(f"/api/v1/teams/{team['id']}/applications", headers=applicant_headers).json()
    invitation = leader.post(f"/api/v1/teams/{team['id']}/invitations", headers=leader_headers, json={"student_id": other_id})
    assert invitation.status_code == 201, invitation.text
    decision = leader.post(f"/api/v1/team-requests/{application['id']}/decision?decision=APPROVED", headers=leader_headers)
    assert decision.status_code == 200, decision.text
    accepted = other.post(f"/api/v1/team-requests/{invitation.json()['id']}/respond?decision=APPROVED", headers=other_headers)
    assert accepted.status_code == 200, accepted.text
    assert leader.post(url, headers=leader_headers).status_code == 200
    detail = leader.get(f"/api/v1/teams/{team['id']}").json()
    assert detail["member_count"] == 3 and detail["open_recruitment"] is False
    processed = leader.get(f"/api/v1/team-requests?class_id={course['id']}").json()["items"]
    assert sum(item["status"] == "APPROVED" for item in processed) == 2
    with SessionLocal() as db:
        records = db.scalars(select(TeamRequest).where(TeamRequest.team_id == UUID(team["id"]))).all()
        assert all(record.resolved_at is not None for record in records if record.status == "CANCELLED")
    deadline = teacher.patch(f"/api/v1/classes/{course['id']}", headers=headers, json={"team_deadline": "2020-01-01T00:00:00+08:00", "version": course["version"]})
    assert deadline.status_code == 200, deadline.text
    expired = leader.post(open_url, headers=leader_headers)
    assert expired.status_code == 409 and expired.json()["code"] == "TEAM_DEADLINE_PASSED"
    teacher.patch(f"/api/v1/classes/{course['id']}", headers=headers, json={"status": "ARCHIVED", "version": deadline.json()["version"]})
    assert leader.post(url, headers=leader_headers).status_code == 409
    assert leader.post(open_url, headers=leader_headers).status_code == 409


def test_password_changes_revoke_existing_sessions():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    second_teacher, _ = login("teacher", "123456", "teacher")

    changed = teacher.post("/api/v1/auth/password", headers=teacher_headers, json={"current_password": "123456", "new_password": "new-password-123"})
    assert changed.status_code == 204, changed.text
    assert teacher.get("/api/v1/auth/session").status_code == 401
    assert second_teacher.get("/api/v1/auth/session").status_code == 401

    teacher, teacher_headers = login("teacher", "new-password-123", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "会话测试", "name": "密码重置班"}).json()
    student = teacher.post(f"/api/v1/classes/{course['id']}/members", headers=teacher_headers, json={"student_no": "20990001", "name": "会话测试学生"}).json()
    student_client, _ = login("20990001", "20990001", "student")
    reset = teacher.post(f"/api/v1/classes/{course['id']}/members/{student['id']}/reset-password", headers=teacher_headers)
    assert reset.status_code == 204, reset.text
    assert student_client.get("/api/v1/auth/session").status_code == 401


def test_concurrent_submission_with_same_idempotency_key_creates_one_version():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "并发测试", "name": "并发提交班"}).json()
    student = teacher.post(f"/api/v1/classes/{course['id']}/members", headers=teacher_headers, json={"student_no": "20990002", "name": "并发测试学生"}).json()
    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": course["id"], "title": "并发提交作业", "description": "验证幂等提交", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    first, first_headers = login(student["student_no"], student["student_no"], "student")
    second, second_headers = login(student["student_no"], student["student_no"], "student")
    team = first.post("/api/v1/teams", headers=first_headers, json={"class_id": course["id"], "name": "并发测试组", "open_recruitment": False})
    assert team.status_code == 201, team.text
    uploaded = first.post(f"/api/v1/assignments/{assignment['id']}/files", headers=first_headers, files={"file": ("work.pdf", b"work", "application/pdf")})
    assert uploaded.status_code == 201, uploaded.text
    barrier = Barrier(2)

    def submit(client, headers):
        barrier.wait()
        return client.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**headers, "Idempotency-Key": "concurrent-submit"}, json={})

    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda args: submit(*args), [(first, first_headers), (second, second_headers)]))

    assert [response.status_code for response in responses] == [201, 201]
    assert responses[0].json()["id"] == responses[1].json()["id"]
    with SessionLocal() as db:
        submission = db.scalar(select(Submission).where(Submission.assignment_id == UUID(assignment["id"]), Submission.owner_user_id == UUID(student["id"])))
        versions = db.scalars(select(SubmissionVersion).where(SubmissionVersion.submission_id == submission.id)).all()
        assert submission.current_version_no == 1
        assert len(versions) == 1


def test_audit_log_search_and_operator_ip():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    created = teacher.post(
        "/api/v1/classes",
        headers=teacher_headers,
        json={"semester": "审计测试", "name": "审计查询班", "max_team_members": 5},
    )
    assert created.status_code == 201, created.text

    result = teacher.get(
        "/api/v1/audit-logs",
        params={"q": "testclient", "actions": "CLASS_CREATED", "page": 1, "page_size": 10},
    )
    assert result.status_code == 200, result.text
    payload = result.json()
    assert payload["total"] >= 1
    assert any(item["action"] == "CLASS_CREATED" and item["ip_address"] == "testclient" and item["class_id"] == created.json()["id"] and item["class_semester"] == "审计测试" and item["class_name"] == "审计查询班" for item in payload["items"])
    filtered = teacher.get("/api/v1/audit-logs", params={"semester": "审计测试", "class_id": created.json()["id"], "actor_role": "TEACHER"})
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()["total"] >= 1
    assert all(item["class_id"] == created.json()["id"] and item["actor_role"] == "TEACHER" for item in filtered.json()["items"])
    deleted = teacher.delete(f"/api/v1/classes/{created.json()['id']}", headers=teacher_headers)
    assert deleted.status_code == 204, deleted.text


def test_formal_course_workflow():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    assert teacher.get("/api/v1/classes").json()["total"] == 0
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2026 秋季", "name": "软件工程 1 班", "max_team_members": 5}).json()
    class_id = course["id"]
    roster = "学号,姓名\n20260001,张同学\n20260002,李同学\n"
    preview = teacher.post(f"/api/v1/classes/{class_id}/members/import-preview", headers=teacher_headers, files={"file": ("名单.csv", roster.encode("utf-8"), "text/csv")})
    assert preview.status_code == 201, preview.text
    assert preview.json()["summary"]["READY"] == 2
    result = teacher.post(f"/api/v1/classes/{class_id}/members/import/{preview.json()['batch_id']}/confirm", headers=teacher_headers)
    assert result.json() == {"created": 2, "joined": 2, "skipped": 0}
    exported_result = teacher.get(f"/api/v1/classes/{class_id}/members/import/{preview.json()['batch_id']}/result.csv")
    assert exported_result.status_code == 200 and "20260001" in exported_result.text

    leader, leader_headers = login("20260001", "20260001", "student")
    applicant, applicant_headers = login("20260002", "20260002", "student")
    added = teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": "20260003", "name": "王同学"})
    assert added.status_code == 201, added.text
    other_leader, other_leader_headers = login("20260003", "20260003", "student")
    assert applicant.get("/api/v1/classes/current/context").json()["team_gate_required"] is True
    team = leader.post("/api/v1/teams", headers=leader_headers, json={"class_id": class_id, "name": "第一小组", "open_recruitment": True})
    assert team.status_code == 201, team.text
    other_team = other_leader.post("/api/v1/teams", headers=other_leader_headers, json={"class_id": class_id, "name": "第二小组", "open_recruitment": True})
    assert other_team.status_code == 201, other_team.text
    empty_export = teacher.get(f"/api/v1/teams/{team.json()['id']}/coursework.zip")
    assert empty_export.status_code == 409 and empty_export.json()["code"] == "NO_COURSEWORK_TO_EXPORT"
    assert [item["id"] for item in leader.get(f"/api/v1/teams?class_id={class_id}").json()["items"]] == [team.json()["id"]]
    assert [item["id"] for item in other_leader.get(f"/api/v1/teams?class_id={class_id}").json()["items"]] == [other_team.json()["id"]]
    assert leader.get(f"/api/v1/teams/{other_team.json()['id']}").status_code == 403
    assert len(applicant.get(f"/api/v1/teams?class_id={class_id}").json()["items"]) == 2
    application = applicant.post(f"/api/v1/teams/{team.json()['id']}/applications", headers=applicant_headers)
    assert application.status_code == 201, application.text
    request_id = leader.get(f"/api/v1/team-requests?class_id={class_id}").json()["items"][0]["id"]
    decision = leader.post(f"/api/v1/team-requests/{request_id}/decision?decision=APPROVED", headers=leader_headers)
    assert decision.status_code == 200, decision.text
    assert applicant.get("/api/v1/classes/current/context").json()["team_gate_required"] is False
    assert [item["id"] for item in applicant.get(f"/api/v1/teams?class_id={class_id}").json()["items"]] == [team.json()["id"]]

    team_assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "小组设计稿", "description": "协作提交设计稿", "submitter_type": "TEAM", "due_at": "2027-12-01T12:00:00+08:00", "allow_late": False, "publish": False})
    assert team_assignment.status_code == 201 and team_assignment.json()["status"] == "DRAFT"
    team_assignment_id = team_assignment.json()["id"]
    teacher_file = teacher.post(f"/api/v1/assignments/{team_assignment_id}/files", headers=teacher_headers, files={"file": ("guide.md", b"# guide", "text/markdown")})
    assert teacher_file.status_code == 201, teacher_file.text
    assert teacher.post(f"/api/v1/assignments/{team_assignment_id}/publish", headers=teacher_headers).json()["status"] == "PUBLISHED"
    forbidden_file = applicant.delete(f"/api/v1/files/{teacher_file.json()['id']}", headers=applicant_headers)
    assert forbidden_file.status_code == 403 and forbidden_file.json()["code"] == "FILE_FORBIDDEN"
    removed_file = teacher.delete(f"/api/v1/files/{teacher_file.json()['id']}", headers=teacher_headers)
    assert removed_file.status_code == 204
    assert applicant.get(f"/api/v1/assignments/{team_assignment_id}/files").json()["attachments"] == []
    assert applicant.get(f"/api/v1/files/{teacher_file.json()['id']}").status_code == 404
    teammate_file = applicant.post(f"/api/v1/assignments/{team_assignment_id}/files", headers=applicant_headers, files={"file": ("design.pdf", b"team draft", "application/pdf")}).json()
    leader_drafts = leader.get(f"/api/v1/assignments/{team_assignment_id}/files").json()["drafts"]
    assert leader_drafts[0]["owner_name"] == "李同学"
    teammate_submit = applicant.post(f"/api/v1/assignments/{team_assignment_id}/submission", headers={**applicant_headers, "Idempotency-Key": "teammate-submit"}, json={})
    assert teammate_submit.status_code == 403 and teammate_submit.json()["code"] == "TEAM_LEADER_REQUIRED"
    team_submit = leader.post(f"/api/v1/assignments/{team_assignment_id}/submission", headers={**leader_headers, "Idempotency-Key": "team-submit"}, json={"file_ids": [teammate_file["id"]]})
    assert team_submit.status_code == 201, team_submit.text
    team_board = teacher.get(f"/api/v1/assignments/{team_assignment_id}/submissions").json()["items"]
    team_version_id = next(item["submission_version_id"] for item in team_board if item["team_id"] == team.json()["id"])
    team_feedback = teacher.post(f"/api/v1/submission-versions/{team_version_id}/feedback/publish", headers=teacher_headers, json={"revision": 0, "grade": "A", "comment": "<p>小组完成度高</p>", "annotations": []})
    assert team_feedback.status_code == 200, team_feedback.text
    graded_team = next(item for item in teacher.get(f"/api/v1/assignments/{team_assignment_id}/submissions").json()["items"] if item["team_id"] == team.json()["id"])
    assert graded_team["teacher_grade"]["grade"] == "A" and graded_team["grading_status"] == "GRADED"
    assert applicant.get(f"/api/v1/submission-versions/{team_version_id}/feedback").json()["grade"] == "A"
    cleared_team = teacher.delete(f"/api/v1/submission-versions/{team_version_id}/feedback", headers=teacher_headers)
    assert cleared_team.status_code == 200 and cleared_team.json()["teacher_grade"] is None
    regraded_team = teacher.post(f"/api/v1/submission-versions/{team_version_id}/feedback/publish", headers=teacher_headers, json={"revision": 0, "grade": "A", "comment": "<p>归档成绩</p>", "annotations": []})
    assert regraded_team.status_code == 200, regraded_team.text
    assert leader.get(f"/api/v1/teams/{team.json()['id']}/coursework.zip").status_code == 403
    assert teacher.get(f"/api/v1/teams/{other_team.json()['id']}/coursework.zip").status_code == 200
    team_export = teacher.get(f"/api/v1/teams/{team.json()['id']}/coursework.zip")
    assert team_export.status_code == 200
    with ZipFile(BytesIO(team_export.content)) as archive:
        names = archive.namelist()
        assert "小组作业提交记录.csv" in names and "小组作业成绩表.csv" in names
        assert any(name.endswith("design.pdf") for name in names)
        assert "小组设计稿" in archive.read("小组作业提交记录.csv").decode("utf-8-sig")
        grade_sheet = archive.read("小组作业成绩表.csv").decode("utf-8-sig")
        assert "学生互评等级,教师等级,最终等级,成绩来源" in grade_sheet
        assert "小组设计稿,已提交,,A,A,教师评分,已评分" in grade_sheet

    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "个人需求报告", "description": "提交个人需求分析作品", "submitter_type": "INDIVIDUAL", "due_at": "2027-12-01T12:00:00+08:00", "allow_late": False, "publish": True})
    assert assignment.status_code == 201, assignment.text
    assignment_id = assignment.json()["id"]
    personal_files = []
    for student, headers, content in [(leader, leader_headers, b"leader work"), (applicant, applicant_headers, b"applicant work")]:
        upload = student.post(f"/api/v1/assignments/{assignment_id}/files", headers=headers, files={"file": ("work.pdf", content, "application/pdf")})
        assert upload.status_code == 201, upload.text
        personal_files.append(upload.json()["id"])
        submitted = student.post(f"/api/v1/assignments/{assignment_id}/submission", headers={**headers, "Idempotency-Key": f"submit-{content.decode()}"}, json={"file_ids": [upload.json()["id"]]})
        assert submitted.status_code == 201, submitted.text

    assert leader.get(f"/api/v1/files/{personal_files[1]}").status_code == 200

    assert teacher.post(f"/api/v1/assignments/{assignment_id}/close", headers=teacher_headers).status_code == 200
    campaign = teacher.post("/api/v1/review-campaigns", headers=teacher_headers, json={"assignment_id": assignment_id, "mode": "TEAM", "criteria_text": "按作品完整性和表达清晰度评分。", "due_at": "2027-12-10T12:00:00+08:00"})
    assert campaign.status_code == 201, campaign.text
    campaign_id = campaign.json()["id"]
    task = leader.get(f"/api/v1/review-campaigns/{campaign_id}/assignment").json()
    assert task["reviewee"]["name"] == "李同学"
    review = leader.post(f"/api/v1/review-campaigns/{campaign_id}/reviews", headers=leader_headers, json={"score": 86, "comment": "作品结构完整，表达清楚。"})
    assert review.status_code == 201, review.text
    assert review.json()["total_score"] == 86
    assert leader.get(f"/api/v1/files/{personal_files[1]}").status_code == 200
    teacher_reviews = teacher.get(f"/api/v1/review-campaigns/{campaign_id}/reviews").json()["items"]
    assert teacher_reviews[0]["reviewee_name"] == "李同学"

    closed_campaign = teacher.post(f"/api/v1/review-campaigns/{campaign_id}/close", headers=teacher_headers)
    assert closed_campaign.status_code == 200
    assert closed_campaign.json()["status"] == "CLOSED" and closed_campaign.json()["grades_generated_at"] is not None
    assert closed_campaign.json()["assignment_id"] == assignment_id
    received = applicant.get(f"/api/v1/peer-reviews/received?class_id={class_id}").json()["items"]
    assert received[0]["reviewer_name"] == "张同学"
    gradebook = teacher.get(f"/api/v1/grades/assignments/{assignment_id}").json()
    assert gradebook["summary"] == {"publishable": 0, "pending": 1, "changed": 0, "published": 0}
    coefficient = gradebook["groups"][0]
    assert teacher.patch(f"/api/v1/grades/assignments/{assignment_id}/teams/{team.json()['id']}/coefficient", headers=teacher_headers, json={"coefficient": -0.1, "version": coefficient["version"]}).status_code == 422
    assert teacher.patch(f"/api/v1/grades/assignments/{assignment_id}/teams/{team.json()['id']}/coefficient", headers=teacher_headers, json={"coefficient": 1.001, "version": coefficient["version"]}).status_code == 422
    saved = teacher.patch(f"/api/v1/grades/assignments/{assignment_id}/teams/{team.json()['id']}/coefficient", headers=teacher_headers, json={"coefficient": 1, "version": coefficient["version"]})
    assert saved.status_code == 200, saved.text
    published = teacher.post(f"/api/v1/grades/assignments/{assignment_id}/publish", headers=teacher_headers, json={"reason": ""})
    assert published.status_code == 200 and published.json()["published"] == 1 and published.json()["pending"] == 1
    student_grade = applicant.get(f"/api/v1/grades?class_id={class_id}").json()["items"][0]
    assert float(student_grade["peer_score"]) == 86 and float(student_grade["score"]) == 86

    refreshed = teacher.get(f"/api/v1/grades/assignments/{assignment_id}").json()
    coefficient = refreshed["groups"][0]
    teacher.patch(f"/api/v1/grades/assignments/{assignment_id}/teams/{team.json()['id']}/coefficient", headers=teacher_headers, json={"coefficient": 0.95, "version": coefficient["version"]})
    assert float(applicant.get(f"/api/v1/grades?class_id={class_id}").json()["items"][0]["score"]) == 86
    no_reason = teacher.post(f"/api/v1/grades/assignments/{assignment_id}/publish", headers=teacher_headers, json={"reason": ""})
    assert no_reason.status_code == 422
    changed = teacher.post(f"/api/v1/grades/assignments/{assignment_id}/publish", headers=teacher_headers, json={"reason": "复核后调整小组系数"})
    assert changed.status_code == 200 and changed.json()["changed"] == 1
    grade_id = next(item["id"] for item in teacher.get(f"/api/v1/grades/assignments/{assignment_id}").json()["items"] if item["student_name"] == "李同学")
    assert len(teacher.get(f"/api/v1/grades/{grade_id}/revisions").json()["items"]) == 1
    assert float(applicant.get(f"/api/v1/grades?class_id={class_id}").json()["items"][0]["score"]) == 81.7
    missing_assignment = teacher.get(f"/api/v1/exports/grades.csv?class_id={class_id}")
    assert missing_assignment.status_code == 422 and missing_assignment.json()["code"] == "ASSIGNMENT_REQUIRED"
    exported_grades = teacher.get(f"/api/v1/exports/grades.csv?class_id={class_id}&assignment_id={assignment_id}")
    assert exported_grades.status_code == 200 and "互评分" in exported_grades.text and "待处理" in exported_grades.text
    assert teacher.get(f"/api/v1/exports/grades.xlsx?class_id={class_id}&assignment_id={assignment_id}").status_code == 200

    archived = teacher.patch(f"/api/v1/classes/{class_id}", headers=teacher_headers, json={"status": "ARCHIVED", "version": course["version"]})
    assert archived.status_code == 200
    blocked = applicant.post(f"/api/v1/assignments/{assignment_id}/files", headers=applicant_headers, files={"file": ("late.pdf", b"late", "application/pdf")})
    assert blocked.status_code == 409 and blocked.json()["code"] == "CLASS_ARCHIVED"


def test_student_bulk_download_assignment_materials():
    teacher, headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2036 秋季", "name": "资料批量下载测试班"}).json()
    teacher.post(f"/api/v1/classes/{course['id']}/members", headers=headers, json={"student_no": "20369901", "name": "资料下载学生"})
    student, student_headers = login("20369901", "20369901", "student")
    assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": course["id"], "title": "批量下载资料", "description": "下载资料", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": False}).json()
    files = []
    for content in (b"first", b"second"):
        uploaded = teacher.post(f"/api/v1/assignments/{assignment['id']}/files", headers=headers, files={"file": ("guide.md", content, "text/markdown")})
        assert uploaded.status_code == 201, uploaded.text
        files.append(uploaded.json()["id"])
    endpoint = f"/api/v1/assignments/{assignment['id']}/materials.zip"
    params = [("file_ids", fid) for fid in files]
    assert student.get(endpoint, params=params).status_code == 404
    assert teacher.post(f"/api/v1/assignments/{assignment['id']}/publish", headers=headers).status_code == 200
    downloaded = student.get(endpoint, params=params)
    assert downloaded.status_code == 200 and downloaded.headers["content-type"] == "application/zip"
    with ZipFile(BytesIO(downloaded.content)) as archive:
        assert archive.namelist() == ["guide.md", "guide (1).md"]
        assert archive.read("guide.md") == b"first" and archive.read("guide (1).md") == b"second"
    with ZipFile(BytesIO(student.get(endpoint, params={"file_ids": files[0]}).content)) as archive:
        assert archive.namelist() == ["guide.md"]
    criteria = teacher.post(f"/api/v1/assignments/{assignment['id']}/files?purpose=REVIEW_CRITERIA", headers=headers, files={"file": ("criteria.md", b"criteria", "text/markdown")}).json()
    assert student.get(endpoint, params={"file_ids": criteria["id"]}).status_code == 422
    other_assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": course["id"], "title": "另一份作业", "description": "不能混用附件", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    assert student.get(f"/api/v1/assignments/{other_assignment['id']}/materials.zip", params=params).status_code == 422
    assert student.get(endpoint).status_code == 422
    assert TestClient(app).get(endpoint, params=params).status_code == 401
    outsider_class = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2036 秋季", "name": "无权资料班"}).json()
    outsider_assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": outsider_class["id"], "title": "无权访问作业", "description": "权限检查", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    assert student.get(f"/api/v1/assignments/{outsider_assignment['id']}/materials.zip", params=params).status_code == 404
    assert teacher.delete(f"/api/v1/files/{files[0]}", headers=headers).status_code == 204
    assert student.get(endpoint, params=params).status_code == 422


def test_role_mismatch_and_unauthenticated():
    client = TestClient(app)
    assert client.get("/api/v1/classes").status_code == 401
    assert client.post("/api/v1/auth/login", json={"account": "teacher", "password": "123456", "role": "student"}).status_code == 403


def test_peer_grade_decimal_calculation_caps_and_rounds_half_up():
    assert final_score(Decimal("88.88"), Decimal("0.95")) == Decimal("84.44")
    assert final_score(Decimal("95"), Decimal("1.20")) == Decimal("100.00")


def test_assignment_time_window_update_and_review_visibility():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2030 春季", "name": "时间规则测试班", "max_team_members": 5}).json()
    class_id = course["id"]
    student_record = teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": "20300001", "name": "时间测试学生"}).json()
    student, student_headers = login("20300001", "20300001", "student")
    team = student.post("/api/v1/teams", headers=student_headers, json={"class_id": class_id, "name": "时间测试小组", "open_recruitment": True})
    assert team.status_code == 201, team.text

    invalid = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "无效时间", "description": "开始时间不能晚于截止时间", "submitter_type": "INDIVIDUAL", "starts_at": "2099-12-02T12:00:00+08:00", "due_at": "2099-12-01T12:00:00+08:00"})
    assert invalid.status_code == 422 and invalid.json()["code"] == "ASSIGNMENT_TIME_INVALID"

    expired_team = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "过期小组作业", "description": "截止时间必须在未来", "submitter_type": "TEAM", "due_at": "2020-01-01T00:00:00+08:00", "publish": True})
    assert expired_team.status_code == 422 and expired_team.json()["code"] == "TEAM_ASSIGNMENT_DUE_INVALID"
    expired_team_draft = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "过期小组草稿", "description": "草稿可以保存但不能发布", "submitter_type": "TEAM", "due_at": "2020-01-01T00:00:00+08:00", "publish": False})
    assert expired_team_draft.status_code == 201
    blocked_publish = teacher.post(f"/api/v1/assignments/{expired_team_draft.json()['id']}/publish", headers=teacher_headers)
    assert blocked_publish.status_code == 422 and blocked_publish.json()["code"] == "TEAM_ASSIGNMENT_DUE_INVALID"

    future = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "未开始作业", "description": "测试开始时间限制", "submitter_type": "INDIVIDUAL", "starts_at": "2099-12-01T12:00:00+08:00", "due_at": "2099-12-02T12:00:00+08:00"})
    assert future.status_code == 201, future.text
    teacher_assignments = teacher.get(f"/api/v1/assignments?class_id={class_id}").json()["items"]
    assert future.json()["id"] in {item["id"] for item in teacher_assignments}
    assert student.get(f"/api/v1/assignments?class_id={class_id}").json()["items"] == []
    student_dashboard = student.get(f"/api/v1/classes/{class_id}/dashboard").json()
    assert student_dashboard["summary"]["active_assignments"] == 0
    assert student_dashboard["summary"]["submission_assignment_title"] is None
    assert student_dashboard["assignment_history"] == []
    blocked_upload = student.post(f"/api/v1/assignments/{future.json()['id']}/files", headers=student_headers, files={"file": ("future.pdf", b"future", "application/pdf")})
    assert blocked_upload.status_code == 409 and blocked_upload.json()["code"] == "ASSIGNMENT_NOT_STARTED"

    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "可更新作业", "description": "测试截止前更新提交", "submitter_type": "INDIVIDUAL", "due_at": "2099-12-02T12:00:00+08:00"}).json()
    uploaded = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("work.pdf", b"work", "application/pdf")})
    assert uploaded.status_code == 201, uploaded.text
    notes = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("notes.md", b"notes", "text/markdown")})
    submitted = student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "time-window-submit"}, json={"file_ids": [uploaded.json()["id"], notes.json()["id"]]})
    assert submitted.status_code == 201, submitted.text
    removed = student.delete(f"/api/v1/files/{notes.json()['id']}", headers=student_headers)
    assert removed.status_code == 204, removed.text
    assert {file["name"] for file in student.get(f"/api/v1/assignments/{assignment['id']}/files").json()["drafts"]} == {"work.pdf"}
    current = student.get(f"/api/v1/assignments/{assignment['id']}/submission").json()
    assert {file["name"] for file in current["files"]} == {"work.pdf", "notes.md"}
    updated = student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "time-window-update"}, json={"file_ids": [uploaded.json()["id"]]})
    assert updated.status_code == 201
    assert student.post(f"/api/v1/assignments/{assignment['id']}/submission/retract", headers=student_headers).status_code == 404
    board = teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions", headers=teacher_headers).json()["items"]
    assert board[0]["status"] == "SUBMITTED" and board[0]["student_no"] == "20300001"
    assert board[0]["team_name"] == "时间测试小组"
    assert {file["name"] for file in board[0]["files"]} == {"work.pdf"}
    assert "versions" not in board[0] and "version_no" not in board[0]
    with SessionLocal() as db:
        stored = db.scalar(select(Submission).where(Submission.assignment_id == UUID(assignment["id"])))
        assert len(db.scalars(select(SubmissionVersion).where(SubmissionVersion.submission_id == stored.id)).all()) == 1
    dashboard_response = teacher.get(f"/api/v1/classes/{class_id}/dashboard").json()
    dashboard = dashboard_response["summary"]
    assert dashboard["submission_assignment_title"] == "可更新作业" and dashboard["submission_rate"] == 100
    assert dashboard["ungrouped_member_count"] == 0
    assert dashboard["latest_submission"]["assignment_id"] == assignment["id"]
    assert dashboard["latest_submission"]["submission_version_id"] == board[0]["submission_version_id"]
    assert dashboard["latest_submission"]["owner"] == "时间测试学生"
    assignment_history = next(item for item in dashboard_response["assignment_history"] if item["id"] == assignment["id"])
    assert assignment_history["title"] == "可更新作业"
    assert assignment_history["completion_rate"] == 100

    late_assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "已截止作业", "description": "测试截止后不可撤回", "submitter_type": "INDIVIDUAL", "due_at": "2020-12-02T12:00:00+08:00", "allow_late": True}).json()
    late_file = student.post(f"/api/v1/assignments/{late_assignment['id']}/files", headers=student_headers, files={"file": ("late.pdf", b"late", "application/pdf")}).json()
    late_submit = student.post(f"/api/v1/assignments/{late_assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "late-submit"}, json={"file_ids": [late_file["id"]]})
    assert late_submit.status_code == 201, late_submit.text


def test_logout_returns_no_content_and_revokes_session():
    client, headers = login("teacher", "123456", "teacher")
    response = client.post("/api/v1/auth/logout", headers=headers)
    assert response.status_code == 204
    assert client.get("/api/v1/classes").status_code == 401


def test_invite_code_join_and_class_visibility_settings():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    seed = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2028 春季", "name": "邀请码账号班", "max_team_members": 5}).json()
    added = teacher.post(f"/api/v1/classes/{seed['id']}/members", headers=teacher_headers, json={"student_no": "20280001", "name": "邀请码学生"})
    assert added.status_code == 201, added.text
    student, student_headers = login("20280001", "20280001", "student")

    auto = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2028 春季", "name": "自动加入班", "max_team_members": 5, "topic_public": True, "invite_requires_approval": False}).json()
    joined = student.post("/api/v1/classes/join", headers=student_headers, json={"invite_code": auto["invite_code"]})
    assert joined.status_code == 200 and joined.json()["status"] == "APPROVED"
    assert student.get("/api/v1/classes").json()["total"] >= 2

    approval = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2028 春季", "name": "审核加入班", "max_team_members": 5}).json()
    pending = student.post("/api/v1/classes/join", headers=student_headers, json={"invite_code": approval["invite_code"]})
    assert pending.status_code == 200 and pending.json()["status"] == "PENDING"
    requests = teacher.get(f"/api/v1/classes/{approval['id']}/join-requests").json()["items"]
    assert len(requests) == 1 and requests[0]["student_no"] == "20280001"
    decided = teacher.post(f"/api/v1/classes/{approval['id']}/join-requests/{requests[0]['id']}/decision?decision=APPROVED", headers=teacher_headers)
    assert decided.status_code == 200 and decided.json()["status"] == "APPROVED"

    updated = teacher.patch(f"/api/v1/classes/{approval['id']}", headers=teacher_headers, json={"version": approval["version"], "topic_public": True, "invite_requires_approval": False})
    assert updated.status_code == 200, updated.text
    assert updated.json()["topic_public"] is True and updated.json()["invite_requires_approval"] is False


def test_teacher_can_approve_pending_topic():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2028 秋季", "name": "选题审核测试班", "max_team_members": 5}).json()
    member = teacher.post(f"/api/v1/classes/{course['id']}/members", headers=teacher_headers, json={"student_no": "20280002", "name": "选题学生"})
    assert member.status_code == 201, member.text
    student, student_headers = login("20280002", "20280002", "student")
    team = student.post("/api/v1/teams", headers=student_headers, json={"class_id": course["id"], "name": "选题小组", "open_recruitment": True})
    assert team.status_code == 201, team.text
    reminders = student.get("/api/v1/notifications").json()["items"]
    topic_reminder = next(item for item in reminders if item["kind"] == "TOPIC_REQUIRED")
    assert topic_reminder["link"] == f"/teams?team={team.json()['id']}"
    topic = student.post(f"/api/v1/teams/{team.json()['id']}/topic", headers=student_headers, json={"name": "课程作业系统", "description": "完成课程作业的协作系统"})
    assert topic.status_code == 200, topic.text
    assert all(item["kind"] != "TOPIC_REQUIRED" for item in student.get("/api/v1/notifications").json()["items"])
    decision = teacher.post(f"/api/v1/topics/{topic.json()['id']}/decision?decision=APPROVED", headers=teacher_headers, json={"reason": "审核通过"})
    assert decision.status_code == 200, decision.text
    assert decision.json()["status"] == "APPROVED"


def test_multisheet_roster_and_member_crud():
    teacher, headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2026 秋季", "name": "成员管理测试班", "max_team_members": 5})
    assert course.status_code == 201, course.text
    class_id = course.json()["id"]

    created = teacher.post(f"/api/v1/classes/{class_id}/members", headers=headers, json={"student_no": "20269991", "name": "新增学生"})
    assert created.status_code == 201, created.text
    member_id = created.json()["id"]
    assert teacher.get(f"/api/v1/classes/{class_id}/members/{member_id}").json()["student_no"] == "20269991"
    duplicate = teacher.post(f"/api/v1/classes/{class_id}/members", headers=headers, json={"student_no": "20269991", "name": "重复学生"})
    assert duplicate.status_code == 409 and duplicate.json()["code"] == "MEMBER_EXISTS"

    updated = teacher.patch(f"/api/v1/classes/{class_id}/members/{member_id}", headers=headers, json={"name": "更正姓名"})
    assert updated.status_code == 200 and updated.json()["name"] == "更正姓名"
    removed = teacher.delete(f"/api/v1/classes/{class_id}/members/{member_id}", headers=headers)
    assert removed.status_code == 204
    assert teacher.get(f"/api/v1/classes/{class_id}/members").json()["total"] == 0
    restored = teacher.post(f"/api/v1/classes/{class_id}/members", headers=headers, json={"student_no": "20269991", "name": "更正姓名"})
    assert restored.status_code == 201 and restored.json()["id"] == member_id

    workbook = Workbook()
    roster = workbook.active
    roster.title = "学生名单"
    roster.append(["学号", "姓名"])
    roster.append([20269992, "表格学生"])
    rooms = workbook.create_sheet("寝室信息")
    rooms.append(["寝室号", "寝室长"])
    rooms.append([501, "表格学生"])
    workbook.active = 1
    content = BytesIO()
    workbook.save(content)

    preview = teacher.post(f"/api/v1/classes/{class_id}/members/import-preview", headers=headers, files={"file": ("多工作表名单.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert preview.status_code == 201, preview.text
    assert preview.json()["summary"]["READY"] == 1
    result = teacher.post(f"/api/v1/classes/{class_id}/members/import/{preview.json()['batch_id']}/confirm", headers=headers)
    assert result.json() == {"created": 1, "joined": 1, "skipped": 0}
    assert teacher.get(f"/api/v1/classes/{class_id}/members").json()["total"] == 2


def test_parse_roster_with_incorrect_worksheet_dimension():
    workbook = Workbook()
    roster = workbook.active
    roster.append(["学号", "姓名"])
    roster.append(["20269993", "范围异常学生"])
    content = BytesIO()
    workbook.save(content)

    malformed = BytesIO()
    with ZipFile(BytesIO(content.getvalue())) as source, ZipFile(malformed, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<dimension ref="A1:B2"/>', b'<dimension ref="A1"/>')
            target.writestr(item, data)

    assert parse_roster(malformed.getvalue(), "范围异常名单.xlsx") == [("20269993", "范围异常学生")]


def test_classes_are_ordered_by_semester_name_and_creation_time():
    teacher, headers = login("teacher", "123456", "teacher")
    oldest_same_name = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2038-2039-1", "name": "同名班"}).json()
    newest_same_name = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2038-2039-1", "name": "同名班"}).json()
    older_semester = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2037-2038-2", "name": "A班"}).json()
    later_name = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2039-2040-1", "name": "B班"}).json()
    earlier_name = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2039-2040-1", "name": "A班"}).json()

    items = teacher.get("/api/v1/classes").json()["items"]
    assert [item["id"] for item in items] == [
        earlier_name["id"], later_name["id"], newest_same_name["id"], oldest_same_name["id"], older_semester["id"],
    ]


def test_class_management_and_multi_class_creation_are_atomic():
    teacher, headers = login("teacher", "123456", "teacher")
    first = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2027 春季", "name": "跨班测试一班"}).json()
    second = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2027 春季", "name": "跨班测试二班", "max_team_members": 5}).json()
    disposable = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2027 春季", "name": "待删除空班", "max_team_members": 5}).json()

    listed = {item["id"]: item for item in teacher.get("/api/v1/classes").json()["items"]}
    assert listed[first["id"]]["member_count"] == 0
    assert listed[first["id"]]["assignment_count"] == 0
    assert listed[first["id"]]["deletable"] is True

    updated = teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"version": first["version"], "semester": "2027 春季", "name": "跨班测试 A 班", "max_team_members": 6, "team_deadline": "2027-03-01T12:00:00+08:00"})
    assert updated.status_code == 200, updated.text
    assert updated.json()["name"] == "跨班测试 A 班" and updated.json()["version"] == 2
    conflict = teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"version": 1, "name": "过期修改"})
    assert conflict.status_code == 409 and conflict.json()["code"] == "CLASS_VERSION_CONFLICT"

    archived = teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"version": 2, "status": "ARCHIVED"}).json()
    blocked_edit = teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"version": archived["version"], "name": "归档后修改"})
    assert blocked_edit.status_code == 409 and blocked_edit.json()["code"] == "CLASS_ARCHIVED"
    restored = teacher.patch(f"/api/v1/classes/{first['id']}", headers=headers, json={"version": archived["version"], "status": "ACTIVE"})
    assert restored.status_code == 200

    assert teacher.delete(f"/api/v1/classes/{disposable['id']}", headers=headers).status_code == 204
    assert disposable["id"] not in {item["id"] for item in teacher.get("/api/v1/classes").json()["items"]}

    student_ids = []
    for index in range(3):
        student_no = f"2027999{index}"
        member = teacher.post(f"/api/v1/classes/{first['id']}/members", headers=headers, json={"student_no": student_no, "name": f"跨班学生{index}"})
        assert member.status_code == 201, member.text
        student_ids.append((student_no, member.json()["id"]))
    ungrouped_summary = teacher.get(f"/api/v1/classes/{first['id']}/dashboard").json()["summary"]
    assert ungrouped_summary["ungrouped_member_count"] == 3
    assert ungrouped_summary["latest_submission"] is None
    leader, leader_headers = login(student_ids[0][0], student_ids[0][0], "student")
    team = leader.post("/api/v1/teams", headers=leader_headers, json={"class_id": first["id"], "name": "跨班测试小组", "open_recruitment": True}).json()
    with SessionLocal() as db:
        stored_team = db.get(Team, UUID(team["id"]))
        stored_team.max_members = 2
        db.commit()
    for student_no, _ in student_ids[1:]:
        applicant, applicant_headers = login(student_no, student_no, "student")
        request = applicant.post(f"/api/v1/teams/{team['id']}/applications", headers=applicant_headers)
        assert request.status_code == 201, request.text
        decision = leader.post(f"/api/v1/team-requests/{request.json()['id']}/decision?decision=APPROVED", headers=leader_headers)
        assert decision.status_code == 200, decision.text
    unlimited_team = leader.get(f"/api/v1/teams/{team['id']}").json()
    assert len(unlimited_team["members"]) == 3

    not_empty = teacher.delete(f"/api/v1/classes/{first['id']}", headers=headers)
    assert not_empty.status_code == 409 and not_empty.json()["code"] == "CLASS_NOT_EMPTY"
    assert not_empty.json()["details"]["members"] == 3

    auto_group_course = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2027 春季", "name": "自动分组测试班"}).json()
    for index in range(13):
        teacher.post(f"/api/v1/classes/{auto_group_course['id']}/members", headers=headers, json={"student_no": f"2027888{index:02d}", "name": f"分组学生{index}"})
    grouped = teacher.post(f"/api/v1/classes/{auto_group_course['id']}/teams/auto-group", headers=headers, json={"group_size": 6})
    assert grouped.status_code == 201, grouped.text
    assert grouped.json()["created"] == 3 and grouped.json()["assigned"] == 13
    assert [item["member_count"] for item in grouped.json()["teams"]] == [6, 6, 1]
    assert teacher.post(f"/api/v1/classes/{auto_group_course['id']}/teams/auto-group", headers=headers, json={"group_size": 6}).json() == {"created": 0, "assigned": 0, "teams": []}

    due_at = "2027-12-01T12:00:00+08:00"
    bulk = teacher.post("/api/v1/assignments/bulk", headers=headers, json={"class_ids": [first["id"], second["id"]], "title": "跨班个人作业", "description": "各班独立完成", "submitter_type": "INDIVIDUAL", "due_at": due_at, "publish": True})
    assert bulk.status_code == 201, bulk.text
    assert bulk.json()["total"] == 2
    created_by_class = {item["class_id"]: item for item in bulk.json()["items"]}
    assert len({item["id"] for item in bulk.json()["items"]}) == 2
    assert teacher.get(f"/api/v1/assignments?class_id={first['id']}").json()["total"] == 1
    assert teacher.get(f"/api/v1/assignments?class_id={second['id']}").json()["total"] == 1

    second_archived = teacher.patch(f"/api/v1/classes/{second['id']}", headers=headers, json={"version": second["version"], "status": "ARCHIVED"}).json()
    failed_bulk = teacher.post("/api/v1/assignments/bulk", headers=headers, json={"class_ids": [first["id"], second["id"]], "title": "不应部分创建", "description": "原子性验证", "submitter_type": "INDIVIDUAL", "due_at": due_at})
    assert failed_bulk.status_code == 409 and failed_bulk.json()["code"] == "CLASS_ARCHIVED"
    assert teacher.get(f"/api/v1/assignments?class_id={first['id']}").json()["total"] == 1
    teacher.patch(f"/api/v1/classes/{second['id']}", headers=headers, json={"version": second_archived["version"], "status": "ACTIVE"})

    removed_legacy_api = teacher.post("/api/v1/review-campaigns/bulk", headers=headers, json={"targets": []})
    assert removed_legacy_api.status_code == 404

    student, student_headers = login(student_ids[1][0], student_ids[1][0], "student")
    forbidden = student.patch(f"/api/v1/classes/{first['id']}", headers=student_headers, json={"version": restored.json()["version"], "name": "越权修改"})
    assert forbidden.status_code == 403


def test_one_to_one_class_review_freezes_assignment_and_validates_score():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2031 秋季", "name": "一对一互评测试班", "max_team_members": 5}).json()
    class_id = course["id"]
    students = []
    for index in range(3):
        student_no = f"2031888{index}"
        teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": student_no, "name": f"互评学生{index}"})
        students.append(login(student_no, student_no, "student"))

    leader, leader_headers = students[0]
    team = leader.post("/api/v1/teams", headers=leader_headers, json={"class_id": class_id, "name": "互评测试组", "open_recruitment": True}).json()
    for student, headers in students[1:]:
        request = student.post(f"/api/v1/teams/{team['id']}/applications", headers=headers).json()
        assert leader.post(f"/api/v1/team-requests/{request['id']}/decision?decision=APPROVED", headers=leader_headers).status_code == 200

    assignment_response = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "冻结版本作业", "description": '<h2>要求</h2><script>alert(1)</script><p><strong>正文</strong></p>', "submitter_type": "INDIVIDUAL", "due_at": "2020-01-01T00:00:00+08:00", "allow_late": True})
    assert assignment_response.status_code == 201, assignment_response.text
    assignment = assignment_response.json()
    assert "<script" not in assignment["description"] and "<strong>正文</strong>" in assignment["description"]

    original_files = {}
    for index, (student, headers) in enumerate(students):
        upload = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=headers, files={"file": (f"work-{index}.pdf", b"pdf", "application/pdf")}).json()
        submitted = student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**headers, "Idempotency-Key": f"frozen-{index}"}, json={"file_ids": [upload["id"]]})
        assert submitted.status_code == 201, submitted.text
        original_files[index] = upload["name"]

    office = teacher.post(f"/api/v1/assignments/{assignment['id']}/files", headers=teacher_headers, files={"file": ("reference.docx", b"docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert office.status_code == 201 and office.json()["download_only"] is True and office.json()["preview_status"] == "NOT_AVAILABLE"
    office_preview = teacher.get(f"/api/v1/files/{office.json()['id']}/preview")
    assert office_preview.status_code == 200 and "attachment" in office_preview.headers["content-disposition"]
    blocked_office = students[0][0].post(f"/api/v1/assignments/{assignment['id']}/files", headers=students[0][1], files={"file": ("work.docx", b"docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert blocked_office.status_code == 422 and blocked_office.json()["code"] == "FILE_TYPE_INVALID"
    assert "学生提交仅支持" in blocked_office.json()["message"]

    campaign = teacher.post("/api/v1/review-campaigns", headers=teacher_headers, json={"assignment_id": assignment["id"], "mode": "TEAM", "criteria_text": "按完整性与清晰度给出总分。", "due_at": "2099-01-01T00:00:00+08:00"})
    assert campaign.status_code == 201, campaign.text
    assert campaign.json()["allocated"] == 3 and campaign.json()["skipped"] == 0
    campaign_id = campaign.json()["id"]

    tasks = [student.get(f"/api/v1/review-campaigns/{campaign_id}/assignment").json() for student, _ in students]
    assert [task["reviewee"]["student_no"] for task in tasks] == ["20318881", "20318882", "20318880"]
    assert [task["submission"]["files"][0]["name"] for task in tasks] == [original_files[1], original_files[2], original_files[0]]

    replacement = students[1][0].post(f"/api/v1/assignments/{assignment['id']}/files", headers=students[1][1], files={"file": ("replacement.pdf", b"new", "application/pdf")}).json()
    assert students[1][0].post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**students[1][1], "Idempotency-Key": "replacement-version"}, json={"file_ids": [replacement["id"]]}).status_code == 201
    assert students[0][0].get(f"/api/v1/review-campaigns/{campaign_id}/assignment").json()["submission"]["files"][0]["name"] == original_files[1]

    blank = students[0][0].post(f"/api/v1/review-campaigns/{campaign_id}/reviews", headers=students[0][1], json={"score": 88, "comment": "   "})
    assert blank.status_code == 422 and blank.json()["code"] == "COMMENT_REQUIRED"
    completed = students[0][0].post(f"/api/v1/review-campaigns/{campaign_id}/reviews", headers=students[0][1], json={"score": 88, "comment": "结构完整。"})
    assert completed.status_code == 201 and completed.json()["total_score"] == 88
    repeated = students[0][0].post(f"/api/v1/review-campaigns/{campaign_id}/reviews", headers=students[0][1], json={"score": 90, "comment": "重复"})
    assert repeated.status_code == 201 and repeated.json()["updated"] is True and repeated.json()["total_score"] == 90
    with SessionLocal() as db:
        reviews = db.scalars(select(PeerReview).where(PeerReview.campaign_id == UUID(campaign_id))).all()
        updates = db.scalars(select(AuditLog).where(AuditLog.action == "PEER_REVIEW_UPDATED", AuditLog.object_id == str(reviews[0].id))).all()
        assert len(reviews) == 1 and reviews[0].comment == "重复"
        assert len(updates) == 1
    stats = teacher.get(f"/api/v1/review-campaigns/{campaign_id}/stats").json()
    assert stats["assigned_count"] == 3 and stats["completed_count"] == 1 and stats["completion_rate"] == 33.3

    locked = teacher.patch(f"/api/v1/assignments/{assignment['id']}", headers=teacher_headers, json={"version": assignment["version"], "auto_review_enabled": False})
    assert locked.status_code == 409 and locked.json()["code"] == "AUTO_REVIEW_CONFIG_LOCKED"

    future = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "尚未截止作业", "description": "不能提前创建互评", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00"}).json()
    blocked = teacher.post("/api/v1/review-campaigns", headers=teacher_headers, json={"assignment_id": future["id"], "mode": "TEAM", "criteria_text": "标准", "due_at": "2099-02-01T00:00:00+08:00"})
    assert blocked.status_code == 409 and blocked.json()["code"] == "ASSIGNMENT_NOT_CLOSED"


def test_team_review_allocates_per_team_and_skips_singleton_team():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2032 春季", "name": "组内分配测试班", "max_team_members": 5}).json()
    class_id = course["id"]
    students = []
    for index in range(3):
        student_no = f"2032777{index}"
        teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": student_no, "name": f"组内学生{index}"})
        students.append(login(student_no, student_no, "student"))
    first_team = students[0][0].post("/api/v1/teams", headers=students[0][1], json={"class_id": class_id, "name": "双人组", "open_recruitment": True}).json()
    request = students[1][0].post(f"/api/v1/teams/{first_team['id']}/applications", headers=students[1][1]).json()
    students[0][0].post(f"/api/v1/team-requests/{request['id']}/decision?decision=APPROVED", headers=students[0][1])
    students[2][0].post("/api/v1/teams", headers=students[2][1], json={"class_id": class_id, "name": "单人组", "open_recruitment": False})

    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "组内互评作业", "description": "组内循环分配", "submitter_type": "INDIVIDUAL", "due_at": "2020-01-01T00:00:00+08:00", "allow_late": True}).json()
    for index, (student, headers) in enumerate(students):
        file = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=headers, files={"file": (f"team-{index}.pdf", b"pdf", "application/pdf")}).json()
        assert student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**headers, "Idempotency-Key": f"team-ring-{index}"}, json={"file_ids": [file["id"]]}).status_code == 201

    campaign = teacher.post("/api/v1/review-campaigns", headers=teacher_headers, json={"assignment_id": assignment["id"], "mode": "TEAM", "criteria_text": "组内互评标准", "due_at": "2099-01-01T00:00:00+08:00"})
    assert campaign.status_code == 201, campaign.text
    assert campaign.json()["allocated"] == 2 and campaign.json()["skipped"] == 1
    assert campaign.json()["warnings"][0]["group"] == "单人组"
    tasks = [student.get(f"/api/v1/review-campaigns/{campaign.json()['id']}/assignment").json() for student, _ in students]
    assert tasks[0]["reviewee"]["student_no"] == "20327771"
    assert tasks[1]["reviewee"]["student_no"] == "20327770"
    assert tasks[2]["status"] == "SKIPPED" and "少于 2 人" in tasks[2]["skip_reason"]


def test_assignment_deadline_automatically_creates_frozen_review():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2033 春季", "name": "自动互评测试班", "max_team_members": 5}).json()
    class_id = course["id"]
    students = []
    for index in range(2):
        student_no = f"2033888{index}"
        teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": student_no, "name": f"自动互评学生{index}"})
        students.append(login(student_no, student_no, "student"))
    team = students[0][0].post("/api/v1/teams", headers=students[0][1], json={"class_id": class_id, "name": "自动互评小组", "open_recruitment": True}).json()
    request = students[1][0].post(f"/api/v1/teams/{team['id']}/applications", headers=students[1][1]).json()
    students[0][0].post(f"/api/v1/team-requests/{request['id']}/decision?decision=APPROVED", headers=students[0][1])

    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={
        "class_id": class_id, "title": "自动互评作业", "description": "截止后自动分配", "submitter_type": "INDIVIDUAL",
        "due_at": "2020-01-01T00:00:00+08:00", "allow_late": True, "auto_review_enabled": True,
        "auto_review_mode": "TEAM", "auto_review_criteria_text": "按完整性评分。", "auto_review_due_at": "2099-01-01T00:00:00+08:00",
    })
    assert assignment.status_code == 201, assignment.text
    assignment_id = assignment.json()["id"]
    for index, (student, headers) in enumerate(students):
        file = student.post(f"/api/v1/assignments/{assignment_id}/files", headers=headers, files={"file": (f"auto-{index}.pdf", b"pdf", "application/pdf")}).json()
        assert student.post(f"/api/v1/assignments/{assignment_id}/submission", headers={**headers, "Idempotency-Key": f"auto-{index}"}, json={"file_ids": [file["id"]]}).status_code == 201

    with SessionLocal.begin() as db:
        process_auto_review(db, datetime.now(UTC))
        process_auto_review(db, datetime.now(UTC))
    with SessionLocal() as db:
        saved_assignment = db.get(Assignment, UUID(assignment_id))
        campaigns = db.scalars(select(ReviewCampaign).where(ReviewCampaign.assignment_id == saved_assignment.id)).all()
        assert len(campaigns) == 1
        campaign = campaigns[0]
        allocations = db.scalars(select(ReviewAssignment).where(ReviewAssignment.campaign_id == campaign.id)).all()
        assert saved_assignment.auto_review_status == "CREATED"
        assert campaign.mode == "TEAM" and campaign.assignment_snapshot_at is not None
        assert len(allocations) == 2 and all(item.status == "PENDING" for item in allocations)
    with SessionLocal.begin() as db:
        campaign = db.scalar(select(ReviewCampaign).where(ReviewCampaign.assignment_id == UUID(assignment_id)))
        campaign.due_at = datetime.now(UTC) - timedelta(seconds=1)
    with SessionLocal.begin() as db:
        process_due_campaign(db, datetime.now(UTC))
        process_due_campaign(db, datetime.now(UTC))
    with SessionLocal() as db:
        campaign = db.scalar(select(ReviewCampaign).where(ReviewCampaign.assignment_id == UUID(assignment_id)))
        generated = db.scalars(select(Grade).where(Grade.assignment_id == UUID(assignment_id))).all()
        assert campaign.status == "CLOSED" and campaign.grades_generated_at is not None
        assert len(generated) == 2 and all(item.peer_score is None and item.status == "PENDING" for item in generated)


def test_peer_review_requires_reviewer_submission_and_submitted_work_is_immediately_gradable():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2034 春季", "name": "即时互评测试班"}).json()
    class_id = course["id"]
    students = []
    for index in range(2):
        student_no = f"2034999{index}"
        teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": student_no, "name": f"即时互评学生{index}"})
        students.append(login(student_no, student_no, "student"))
    team = students[0][0].post("/api/v1/teams", headers=students[0][1], json={"class_id": class_id, "name": "即时互评组", "open_recruitment": True}).json()
    request = students[1][0].post(f"/api/v1/teams/{team['id']}/applications", headers=students[1][1]).json()
    students[0][0].post(f"/api/v1/team-requests/{request['id']}/decision?decision=APPROVED", headers=students[0][1])

    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "即时评价作业", "description": "提交后立即评价", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    attachment = teacher.post(f"/api/v1/assignments/{assignment['id']}/files", headers=teacher_headers, files={"file": ("grading-guide.pdf", b"grading guide", "application/pdf")})
    assert attachment.status_code == 201, attachment.text
    criterion = teacher.post(f"/api/v1/assignments/{assignment['id']}/files?purpose=REVIEW_CRITERIA", headers=teacher_headers, files={"file": ("peer-criteria.md", b"# peer criteria", "text/markdown")})
    assert criterion.status_code == 201, criterion.text
    first_student, first_headers = students[0]
    file = first_student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=first_headers, files={"file": ("first.pdf", b"first", "application/pdf")}).json()
    assert first_student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**first_headers, "Idempotency-Key": "direct-first"}, json={"file_ids": [file["id"]]}).status_code == 201

    reviewer, reviewer_headers = students[1]
    assert reviewer.get(f"/api/v1/peer-review-assignments?class_id={class_id}").json()["items"] == []
    blocked_detail = reviewer.get(f"/api/v1/assignments/{assignment['id']}/peer-review")
    assert blocked_detail.status_code == 409 and blocked_detail.json()["code"] == "REVIEWER_SUBMISSION_REQUIRED"
    first_student_id = first_student.get("/api/v1/auth/session").json()["user"]["id"]
    blocked_simple_review = reviewer.post(f"/api/v1/assignments/{assignment['id']}/peer-reviews", headers=reviewer_headers, json={"reviewee_id": first_student_id, "grade": "A", "comment": "不能提前评价"})
    assert blocked_simple_review.status_code == 409 and blocked_simple_review.json()["code"] == "REVIEWER_SUBMISSION_REQUIRED"
    version_id = next(item["submission_version_id"] for item in teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions").json()["items"] if item["student_no"] == "20349990")
    blocked_feedback = reviewer.get(f"/api/v1/submission-versions/{version_id}/peer-feedback")
    assert blocked_feedback.status_code == 409 and blocked_feedback.json()["code"] == "REVIEWER_SUBMISSION_REQUIRED"
    assert reviewer.get(f"/api/v1/files/{file['id']}").status_code == 403
    assert reviewer.get(f"/api/v1/files/{criterion.json()['id']}").status_code == 403

    reviewer_file = reviewer.post(f"/api/v1/assignments/{assignment['id']}/files", headers=reviewer_headers, files={"file": ("reviewer.pdf", b"reviewer", "application/pdf")}).json()
    assert reviewer.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**reviewer_headers, "Idempotency-Key": "direct-reviewer"}, json={"file_ids": [reviewer_file["id"]]}).status_code == 201
    available = reviewer.get(f"/api/v1/peer-review-assignments?class_id={class_id}").json()["items"]
    assert available[0]["assignment_id"] == assignment["id"] and available[0]["pending_count"] == 1
    assert available[0]["candidates"] == [{"user_id": first_student_id, "name": "即时互评学生0", "student_no": "20349990", "reviewed": False}]
    detail = reviewer.get(f"/api/v1/assignments/{assignment['id']}/peer-review").json()
    assert [item["name"] for item in detail["attachments"]] == ["grading-guide.pdf"]
    assert [item["name"] for item in detail["criteria_files"]] == ["peer-criteria.md"]
    assert reviewer.get(f"/api/v1/files/{detail['attachments'][0]['id']}").status_code == 200
    assert reviewer.get(f"/api/v1/files/{detail['criteria_files'][0]['id']}").status_code == 200
    assert [item["student_no"] for item in detail["candidates"]] == ["20349990"]
    version_id = detail["candidates"][0]["submission_version_id"]
    empty_feedback = reviewer.get(f"/api/v1/submission-versions/{version_id}/peer-feedback").json()
    assert empty_feedback["status"] is None and empty_feedback["grade"] is None
    annotation = {"file_id": file["id"], "kind": "PDF_TEXT_OR_REGION", "mark_type": "HIGHLIGHT", "color": "YELLOW", "anchor": {"page": 1, "rects": [{"x": 0.1, "y": 0.1, "width": 0.2, "height": 0.1}], "quote": ""}, "comment": "<p>这里需要补充</p>"}
    reviewed = reviewer.post(f"/api/v1/submission-versions/{version_id}/peer-feedback/publish", headers=reviewer_headers, json={"revision": 0, "grade": "A", "comment": "<p>完成度高</p>", "annotations": [annotation]})
    assert reviewed.status_code == 200 and reviewed.json()["grade"] == "A" and len(reviewed.json()["annotations"]) == 1
    refreshed_available = reviewer.get(f"/api/v1/peer-review-assignments?class_id={class_id}").json()["items"]
    assert refreshed_available[0]["pending_count"] == 0 and refreshed_available[0]["candidates"][0]["reviewed"] is True
    stale = reviewer.post(f"/api/v1/submission-versions/{version_id}/peer-feedback/publish", headers=reviewer_headers, json={"revision": 0, "grade": "B", "comment": "", "annotations": []})
    assert stale.status_code == 409 and stale.json()["code"] == "FEEDBACK_VERSION_CONFLICT"
    self_review = first_student.get(f"/api/v1/submission-versions/{version_id}/peer-feedback")
    assert self_review.status_code == 422 and self_review.json()["code"] == "SELF_REVIEW_FORBIDDEN"

    board = teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions").json()["items"]
    submitted = next(item for item in board if item["student_no"] == "20349990")
    reviewer_submission = next(item for item in board if item["student_no"] == "20349991")
    assert submitted["teacher_grade"] is None
    reviewer_graded = teacher.post(f"/api/v1/assignments/{assignment['id']}/submissions/{reviewer_submission['user_id']}/grade", headers=teacher_headers, json={"grade": "A", "comment": "互评学生作业完成"})
    assert reviewer_graded.status_code == 201 and reviewer_graded.json()["grade"] == "A"
    graded = teacher.post(f"/api/v1/assignments/{assignment['id']}/submissions/{submitted['user_id']}/grade", headers=teacher_headers, json={"grade": "B", "comment": "需求覆盖完整"})
    assert graded.status_code == 201 and graded.json()["grade"] == "B"
    assert teacher.get(f"/api/v1/classes/{class_id}/dashboard").json()["summary"]["latest_submission"] is None
    refreshed = teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions").json()["items"]
    result = next(item for item in refreshed if item["student_no"] == "20349990")
    assert result["teacher_grade"]["grade"] == "B" and result["peer_grade"] == "A"
    assert result["final_grade"] == "B" and result["grade_source"] == "TEACHER"
    cleared = teacher.delete(f"/api/v1/assignments/{assignment['id']}/submissions/{submitted['user_id']}/grade", headers=teacher_headers)
    assert cleared.status_code == 200 and cleared.json()["final_grade"] == "A" and cleared.json()["grade_source"] == "PEER"
    assert teacher.get(f"/api/v1/classes/{class_id}/dashboard").json()["summary"]["latest_submission"]["submission_version_id"] == version_id
    student_grade = first_student.get(f"/api/v1/grades?class_id={class_id}").json()["items"][0]
    assert student_grade["final_grade"] == "A" and student_grade["grade_source"] == "PEER"
    assert student_grade["peer_feedbacks"][0]["evaluator_name"] == "即时互评学生1"
    assert student_grade["peer_feedbacks"][0]["annotations"][0]["comment"] == "<p>这里需要补充</p>"
    exportable = teacher.get(f"/api/v1/grades/assignments?class_id={class_id}").json()["items"]
    assert exportable[0]["id"] == assignment["id"] and exportable[0]["graded"] == 2 and exportable[0]["total"] == 2
    current_grades = teacher.get(f"/api/v1/exports/grades.csv?class_id={class_id}&assignment_id={assignment['id']}")
    assert current_grades.status_code == 200 and "学生互评等级" in current_grades.text and "成绩来源" in current_grades.text
    assert teacher.post(f"/api/v1/assignments/{assignment['id']}/close", headers=teacher_headers).status_code == 200
    assert teacher.delete(f"/api/v1/files/{attachment.json()['id']}", headers=teacher_headers).status_code == 204
    assert reviewer.get(f"/api/v1/assignments/{assignment['id']}/peer-review").json()["attachments"] == []
    blocked = teacher.post(f"/api/v1/assignments/{assignment['id']}/publish", headers=teacher_headers)
    assert blocked.status_code == 422 and blocked.json()["code"] == "ASSIGNMENT_DUE_INVALID"
    closed = next(item for item in teacher.get(f"/api/v1/assignments?class_id={class_id}").json()["items"] if item["id"] == assignment["id"])
    edited = teacher.patch(f"/api/v1/assignments/{assignment['id']}", headers=teacher_headers, json={"title": "更新后的即时评价作业", "due_at": "2099-01-02T00:00:00+08:00", "version": closed["version"]})
    assert edited.status_code == 200, edited.text
    replacement = teacher.post(f"/api/v1/assignments/{assignment['id']}/files", headers=teacher_headers, files={"file": ("updated-guide.md", b"# updated guide", "text/markdown")})
    assert replacement.status_code == 201, replacement.text
    republished = teacher.post(f"/api/v1/assignments/{assignment['id']}/publish", headers=teacher_headers)
    assert republished.status_code == 200 and republished.json()["status"] == "PUBLISHED"
    repeated = teacher.post(f"/api/v1/assignments/{assignment['id']}/publish", headers=teacher_headers)
    assert repeated.status_code == 200 and repeated.json()["version"] == republished.json()["version"] + 1
    refreshed_detail = reviewer.get(f"/api/v1/assignments/{assignment['id']}/peer-review").json()
    assert refreshed_detail["assignment"]["title"] == "更新后的即时评价作业"
    assert [item["id"] for item in refreshed_detail["attachments"]] == [replacement.json()["id"]]
    assert refreshed_detail["candidates"][0]["submission_version_id"] == version_id
    assert "即时互评学生0" in current_grades.text and "即时互评学生1" in current_grades.text and "已提交" in current_grades.text
    current_reviews = teacher.get(f"/api/v1/exports/reviews.csv?class_id={class_id}")
    assert current_reviews.status_code == 200 and "评价结果" in current_reviews.text and "即时互评学生1" in current_reviews.text and ",A," in current_reviews.text
    assert teacher.get(f"/api/v1/exports/reviews.xlsx?class_id={class_id}").status_code == 200
    members_export = teacher.get(f"/api/v1/exports/members.csv?class_id={class_id}")
    assert members_export.status_code == 200 and "正常" in members_export.text

    overdue = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "截止未交作业", "description": "验证未提交系统等级", "submitter_type": "INDIVIDUAL", "due_at": "2020-01-01T00:00:00+08:00", "allow_late": True, "publish": True}).json()
    overdue_board = teacher.get(f"/api/v1/assignments/{overdue['id']}/submissions").json()["items"]
    missing = next(item for item in overdue_board if item["student_no"] == "20349990")
    assert missing["status"] == "NOT_SUBMITTED" and missing["final_grade"] == "E"
    assert missing["grade_source"] == "SYSTEM" and missing["grading_status"] == "NO_SUBMISSION"
    overdue_grade = next(item for item in first_student.get(f"/api/v1/grades?class_id={class_id}").json()["items"] if item["assignment_id"] == overdue["id"])
    assert overdue_grade["final_grade"] == "E" and overdue_grade["submission_version_id"] is None
    overdue_export = teacher.get(f"/api/v1/exports/grades.csv?class_id={class_id}&assignment_id={overdue['id']}")
    assert "系统判定" in overdue_export.text and ",E," in overdue_export.text
    late_file = first_student.post(f"/api/v1/assignments/{overdue['id']}/files", headers=first_headers, files={"file": ("late.pdf", b"late", "application/pdf")}).json()
    late_submit = first_student.post(f"/api/v1/assignments/{overdue['id']}/submission", headers={**first_headers, "Idempotency-Key": "late-grade-reset"}, json={"file_ids": [late_file["id"]]})
    assert late_submit.status_code == 201
    pending_grade = next(item for item in first_student.get(f"/api/v1/grades?class_id={class_id}").json()["items"] if item["assignment_id"] == overdue["id"])
    assert pending_grade["final_grade"] is None and pending_grade["grade_source"] is None
    assert pending_grade["grading_status"] == "PENDING_ASSESSMENT"

    class_mode = teacher.post("/api/v1/review-campaigns", headers=teacher_headers, json={"assignment_id": assignment["id"], "mode": "CLASS", "criteria_text": "已停用", "due_at": "2099-02-01T00:00:00+08:00"})
    assert class_mode.status_code == 422


def test_rich_preview_feedback_annotations_and_resubmission_history(monkeypatch):
    stored_objects = {}
    monkeypatch.setattr("app.storage.put_object", lambda key, path: stored_objects.__setitem__(key, path.read_bytes()))
    monkeypatch.setattr("app.storage.get_object_bytes", lambda key: stored_objects[key])
    monkeypatch.setattr("app.storage.get_object_stream", lambda key: iter([stored_objects[key]]))
    monkeypatch.setattr("app.storage.delete_object", lambda key: stored_objects.pop(key, None))
    monkeypatch.setattr("app.storage.object_exists", lambda key: key in stored_objects)
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2035 春季", "name": "在线批注测试班"}).json()
    class_id = course["id"]
    teacher.post(f"/api/v1/classes/{class_id}/members", headers=teacher_headers, json={"student_no": "20350001", "name": "批注测试学生"})
    student, student_headers = login("20350001", "20350001", "student")
    student.post("/api/v1/teams", headers=student_headers, json={"class_id": class_id, "name": "批注测试组", "open_recruitment": True})
    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "富文本报告", "description": "验证安全预览和批注", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()

    markdown_content = b"# Heading\n\n- [x] Done\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n![chart](data:image/png;base64,iVBORw0KGgo=)\n\n![unsafe](data:image/svg+xml;base64,PHN2Zz48L3N2Zz4=)"
    markdown_file = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("report.md", markdown_content, "text/markdown")}).json()
    pdf_file = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("diagram.pdf", b"%PDF-1.4", "application/pdf")}).json()
    html_file = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("appendix.html", b'<h2>Safe</h2><script>alert(1)</script><img src="/private.png"><img src="http://localhost:9005/ok.png"><img src="https://example.com/ok.png"><a href="javascript:alert(1)">bad</a>', "text/html")})
    assert pdf_file["download_only"] is True
    assert html_file.status_code == 201 and html_file.json()["download_only"] is True
    submitted = student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "annotation-v1"}, json={})
    assert submitted.status_code == 201, submitted.text
    board_item = teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions").json()["items"][0]
    version_id = board_item["submission_version_id"]

    signed = {}
    def fake_sign_get_url(key, expires, params=None):
        signed.update({"key": key, "expires": expires, "params": params})
        return "https://se-lab.example.test/signed-report.md"
    monkeypatch.setattr("app.storage.sign_get_url", fake_sign_get_url)
    direct_preview = teacher.get(f"/api/v1/files/{markdown_file['id']}/preview-url")
    assert direct_preview.status_code == 200
    assert direct_preview.json() == {"url": "https://se-lab.example.test/signed-report.md", "expires_in": 300}
    assert signed["key"].endswith(".md") and signed["expires"] == 300
    assert signed["params"] == {"response-content-disposition": 'inline; filename="report.md"'}
    rendered_markdown = teacher.get(f"/api/v1/files/{markdown_file['id']}/render")
    assert rendered_markdown.status_code == 410 and rendered_markdown.json()["code"] == "FILE_RENDER_REMOVED"
    rendered_html = teacher.get(f"/api/v1/files/{html_file.json()['id']}/render")
    assert rendered_html.status_code == 410 and rendered_html.json()["code"] == "FILE_RENDER_REMOVED"

    annotation = {"file_id": markdown_file["id"], "kind": "RICH_TEXT_RANGE", "mark_type": "COMMENT", "color": "BLUE", "anchor": {"start": {"block_id": "b0", "offset": 0}, "end": {"block_id": "b0", "offset": 7}, "exact": "Heading", "prefix": "", "suffix": "Done"}, "comment": "<p><strong>重点</strong><script>bad()</script></p>"}
    pure_mark = {"file_id": markdown_file["id"], "kind": "RICH_TEXT_RANGE", "mark_type": "UNDERLINE", "color": "GREEN", "anchor": {"start": {"block_id": "b1", "offset": 0}, "end": {"block_id": "b1", "offset": 4}, "exact": "Done", "prefix": "", "suffix": ""}, "comment": ""}
    legacy_mark = {"file_id": markdown_file["id"], "kind": "RICH_TEXT_RANGE", "anchor": {"start": {"block_id": "b1", "offset": 0}, "end": {"block_id": "b1", "offset": 4}, "exact": "Done", "prefix": "", "suffix": ""}, "comment": ""}
    draft = teacher.put(f"/api/v1/submission-versions/{version_id}/feedback/draft", headers=teacher_headers, json={"revision": 0, "grade": "A", "comment": "<p>总评草稿</p>", "annotations": [annotation, pure_mark, legacy_mark]})
    assert draft.status_code == 200 and draft.json()["status"] == "DRAFT"
    assert student.get(f"/api/v1/submission-versions/{version_id}/feedback").json()["status"] is None
    forbidden = student.put(f"/api/v1/submission-versions/{version_id}/feedback/draft", headers=student_headers, json={"revision": 0, "grade": "A", "comment": "", "annotations": []})
    assert forbidden.status_code == 403

    published = teacher.post(f"/api/v1/submission-versions/{version_id}/feedback/publish", headers=teacher_headers, json={"revision": draft.json()["revision"], "grade": "A", "comment": "<p>已发布总评</p>", "annotations": [annotation, pure_mark, legacy_mark]})
    assert published.status_code == 200 and published.json()["status"] == "PUBLISHED"
    student_feedback = student.get(f"/api/v1/submission-versions/{version_id}/feedback").json()
    assert student_feedback["grade"] == "A" and len(student_feedback["annotations"]) == 3
    assert student_feedback["annotations"][0]["mark_type"] == "COMMENT" and student_feedback["annotations"][0]["color"] == "BLUE"
    assert student_feedback["annotations"][1]["mark_type"] == "UNDERLINE" and student_feedback["annotations"][1]["comment"] == ""
    assert student_feedback["annotations"][2]["mark_type"] == "HIGHLIGHT" and student_feedback["annotations"][2]["color"] == "YELLOW"
    assert "<script" not in student_feedback["annotations"][0]["comment"]
    invalid_region = {"file_id": pdf_file["id"], "kind": "PDF_TEXT_OR_REGION", "mark_type": "UNDERLINE", "color": "RED", "anchor": {"page": 1, "rects": [{"x": 0.1, "y": 0.1, "width": 0.2, "height": 0.1}], "quote": ""}, "comment": ""}
    invalid = teacher.put(f"/api/v1/submission-versions/{version_id}/feedback/draft", headers=teacher_headers, json={"revision": published.json()["revision"], "grade": "A", "comment": "", "annotations": [invalid_region]})
    assert invalid.status_code == 422 and invalid.json()["code"] == "ANNOTATION_MARK_TYPE_INVALID"
    stale = teacher.post(f"/api/v1/submission-versions/{version_id}/feedback/publish", headers=teacher_headers, json={"revision": draft.json()["revision"], "grade": "B", "comment": "", "annotations": []})
    assert stale.status_code == 409 and stale.json()["code"] == "FEEDBACK_VERSION_CONFLICT"

    student.delete(f"/api/v1/files/{markdown_file['id']}", headers=student_headers)
    student.delete(f"/api/v1/files/{pdf_file['id']}", headers=student_headers)
    student.delete(f"/api/v1/files/{html_file.json()['id']}", headers=student_headers)
    replacement = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("report-v2.md", b"# Version 2", "text/markdown")}).json()
    updated = student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "annotation-v2"}, json={})
    assert updated.status_code == 201
    refreshed = teacher.get(f"/api/v1/assignments/{assignment['id']}/submissions").json()["items"][0]
    assert refreshed["submission_version_no"] == 2 and refreshed["submission_version_id"] != version_id
    assert refreshed["teacher_grade"] is None and refreshed["files"][0]["id"] == replacement["id"]
    assert student.get(f"/api/v1/submission-versions/{version_id}/feedback").json()["status"] == "PUBLISHED"


def test_restore_embedded_images_supports_markdown_and_html_without_duplicates():
    from app.main import restore_embedded_images

    markdown_image = "![流程图](data:image/png;base64,aGVsbG8=)"
    html_image = '<img src="data:image/jpeg;base64,d29ybGQ=" alt="截图" width="50%">'
    template = f"# 任务一\n\n{markdown_image}\n\n## 任务二\n\n{html_image}\n"
    draft = "# 任务一\n\n学生答案\n\n## 任务二\n\n补充说明\n"

    restored = restore_embedded_images(template, draft)

    assert markdown_image in restored
    assert html_image in restored
    assert "学生答案" in restored and "补充说明" in restored
    assert restore_embedded_images(template, restored) == restored


def test_restore_embedded_tables_repairs_only_unchanged_flattened_tables():
    from app.main import restore_embedded_tables

    table = "| 我发现的问题 | 属于哪一类 | 怎么改（修法） |\r\n|---|---|---|\r\n|  |  |  |\r\n|  |  |  |"
    template = f"# 任务\r\n\r\n{table}\r\n\r\n---\r\n"
    flattened = "我发现的问题属于哪一类怎么改（修法）"
    draft = f"# 任务\n\n{flattened}\n\n---\n\n学生答案\n"

    restored = restore_embedded_tables(template, draft)

    assert table.replace("\r\n", "\n") in restored
    assert "学生答案" in restored
    assert restore_embedded_tables(template, restored) == restored
    assert restore_embedded_tables(template, draft.replace(flattened, "已修改的表头")) == draft.replace(flattened, "已修改的表头")


def test_workspace_loads_document_content_lazily_and_checks_source_images_once(monkeypatch):
    stored_objects = {}
    reads = []
    monkeypatch.setattr("app.storage.put_object", lambda key, path: stored_objects.__setitem__(key, path.read_bytes()))
    monkeypatch.setattr("app.storage.get_object_bytes", lambda key: reads.append(key) or stored_objects[key])
    monkeypatch.setattr("app.storage.object_exists", lambda key: key in stored_objects)
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2036 春季", "name": "懒加载测试班"}).json()
    teacher.post(f"/api/v1/classes/{course['id']}/members", headers=teacher_headers, json={"student_no": "20360001", "name": "懒加载学生"})
    assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": course["id"], "title": "Markdown 懒加载", "description": "验证目录和正文拆分", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00", "publish": True}).json()
    source = b"# Task\n\n![source](data:image/png;base64,aGVsbG8=)\n"
    uploaded = teacher.post(f"/api/v1/assignments/{assignment['id']}/files?material_type=TASK", headers=teacher_headers, files={"file": ("task.md", source, "text/markdown")})
    assert uploaded.status_code == 201, uploaded.text
    student, student_headers = login("20360001", "20360001", "student")

    workspace = student.post(f"/api/v1/assignments/{assignment['id']}/workspace", headers=student_headers, json={})
    assert workspace.status_code == 200, workspace.text
    metadata = workspace.json()["documents"][0]
    assert "markdown_content" not in metadata and "content_html" not in metadata
    assert len(reads) == 0

    with SessionLocal() as db:
        document = db.get(SubmissionDocument, UUID(metadata["id"]))
        document.markdown_content = "# Student draft\n"
        document.source_images_checked_at = None
        db.commit()

    first = student.get(f"/api/v1/assignments/{assignment['id']}/workspace/documents/{metadata['id']}")
    assert first.status_code == 200, first.text
    assert first.json()["markdown_content"].count("data:image") == 1
    assert "content_html" not in first.json()
    assert len(reads) == 1
    second = student.get(f"/api/v1/assignments/{assignment['id']}/workspace/documents/{metadata['id']}")
    assert second.status_code == 200 and len(reads) == 1

    saved = student.put(f"/api/v1/assignments/{assignment['id']}/workspace/documents/{metadata['id']}", headers=student_headers, json={"revision": first.json()["revision"], "markdown_content": "# Saved\n"})
    assert saved.status_code == 200 and "markdown_content" not in saved.json()
    legacy_payload = student.put(f"/api/v1/assignments/{assignment['id']}/workspace/documents/{metadata['id']}", headers=student_headers, json={"revision": saved.json()["revision"], "markdown_content": "# Saved\n", "content_html": "<h1>Saved</h1>"})
    assert legacy_payload.status_code == 422
    conflict = student.put(f"/api/v1/assignments/{assignment['id']}/workspace/documents/{metadata['id']}", headers=student_headers, json={"revision": first.json()["revision"], "markdown_content": "# Stale\n"})
    assert conflict.status_code == 409
    assert "markdown_content" not in conflict.json()["details"]["document"]


def test_teacher_assignment_and_review_lifecycle_controls():
    teacher, teacher_headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=teacher_headers, json={"semester": "2032 春季", "name": "作业生命周期测试班", "max_team_members": 5}).json()
    class_id = course["id"]

    draft = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "生命周期作业", "description": "可编辑、发布和撤回", "submitter_type": "INDIVIDUAL", "due_at": "2099-12-02T12:00:00+08:00", "publish": False})
    assert draft.status_code == 201 and draft.json()["status"] == "DRAFT"
    assignment_id = draft.json()["id"]
    edited = teacher.patch(f"/api/v1/assignments/{assignment_id}", headers=teacher_headers, json={"title": "已编辑生命周期作业", "version": draft.json()["version"]})
    assert edited.status_code == 200 and edited.json()["title"] == "已编辑生命周期作业"
    configured = teacher.patch(f"/api/v1/assignments/{assignment_id}", headers=teacher_headers, json={
        "version": edited.json()["version"], "auto_review_enabled": True, "auto_review_mode": "TEAM",
        "auto_review_criteria_text": "按完整性评分", "auto_review_due_at": "2099-12-10T12:00:00+08:00",
    })
    assert configured.status_code == 200 and configured.json()["auto_review_enabled"] is True
    reconfigured = teacher.patch(f"/api/v1/assignments/{assignment_id}", headers=teacher_headers, json={
        "version": configured.json()["version"], "auto_review_mode": "TEAM", "auto_review_criteria_text": "按清晰度评分",
        "auto_review_due_at": "2099-12-11T12:00:00+08:00",
    })
    assert reconfigured.status_code == 200 and reconfigured.json()["auto_review_mode"] == "TEAM"
    published = teacher.post(f"/api/v1/assignments/{assignment_id}/publish", headers=teacher_headers)
    assert published.status_code == 200 and published.json()["status"] == "PUBLISHED"
    closed = teacher.post(f"/api/v1/assignments/{assignment_id}/close", headers=teacher_headers)
    assert closed.status_code == 200 and closed.json()["status"] == "CLOSED"
    retracted = teacher.post(f"/api/v1/assignments/{assignment_id}/retract", headers=teacher_headers)
    assert retracted.status_code == 200 and retracted.json()["status"] == "DRAFT"
    assert teacher.delete(f"/api/v1/assignments/{assignment_id}", headers=teacher_headers).status_code == 204

    attachment_draft = teacher.post("/api/v1/assignments", headers=teacher_headers, json={
        "class_id": class_id, "title": "附件标准作业", "description": "标准仅见附件", "submitter_type": "INDIVIDUAL",
        "due_at": "2099-12-02T12:00:00+08:00", "publish": False, "auto_review_enabled": True,
        "auto_review_mode": "TEAM", "auto_review_criteria_text": "", "auto_review_due_at": "2099-12-10T12:00:00+08:00",
    }).json()
    criterion = teacher.post(f"/api/v1/assignments/{attachment_draft['id']}/files?purpose=REVIEW_CRITERIA", headers=teacher_headers, files={"file": ("criteria.pdf", b"pdf", "application/pdf")})
    assert criterion.status_code == 201
    assert teacher.post(f"/api/v1/assignments/{attachment_draft['id']}/publish", headers=teacher_headers).status_code == 200
    last_source = teacher.delete(f"/api/v1/files/{criterion.json()['id']}", headers=teacher_headers)
    assert last_source.status_code == 409 and last_source.json()["code"] == "REVIEW_CRITERIA_REQUIRED"
    assert teacher.delete(f"/api/v1/assignments/{attachment_draft['id']}", headers=teacher_headers).status_code == 204

    review_assignment = teacher.post("/api/v1/assignments", headers=teacher_headers, json={"class_id": class_id, "title": "互评截止作业", "description": "用于互评提前截止", "submitter_type": "INDIVIDUAL", "due_at": "2099-12-02T12:00:00+08:00", "publish": True}).json()
    assert teacher.delete(f"/api/v1/assignments/{review_assignment['id']}", headers=teacher_headers).status_code == 204
    assert teacher.get(f"/api/v1/assignments?class_id={class_id}").json()["items"] == []
