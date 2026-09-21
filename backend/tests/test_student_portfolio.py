from datetime import datetime, timedelta, timezone

UTC = timezone.utc
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.database import SessionLocal
from app.main import app
from app.models import Assignment, BackgroundJob, FileObject, Submission, SubmissionAssessment, SubmissionVersion
from app.student_portfolio import safe_name
from app.worker import process_archive_export
from test_api import login


def test_student_portfolio_and_archives(isolated_object_storage):
    teacher, headers = login("teacher", "123456", "teacher")
    course = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2039 秋季", "name": "档案导出班"}).json()
    cid = course["id"]
    member = teacher.post(f"/api/v1/classes/{cid}/members", headers=headers, json={"student_no": "20390001", "name": "=SUM(1,2)"}).json()
    removed = teacher.post(f"/api/v1/classes/{cid}/members", headers=headers, json={"student_no": "20390002", "name": "移出学生"}).json()
    student, student_headers = login("20390001", "20390001", "student")
    student.post("/api/v1/teams", headers=student_headers, json={"class_id": cid, "name": "档案组"})
    assignment = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": cid, "title": "个人报告", "description": "报告", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-01T00:00:00+08:00"}).json()
    other = teacher.post("/api/v1/assignments", headers=headers, json={"class_id": cid, "title": "未交作业", "description": "报告", "submitter_type": "INDIVIDUAL", "due_at": "2099-01-02T00:00:00+08:00"}).json()
    teacher.post("/api/v1/assignments", headers=headers, json={"class_id": cid, "title": "小组报告", "description": "报告", "submitter_type": "TEAM", "due_at": "2099-01-01T00:00:00+08:00"})
    files = []
    for content in [b"one", b"two"]:
        response = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("report.pdf", content, "application/pdf")})
        assert response.status_code == 201, response.text
        files.append(response.json())
    png = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    markdown = student.post(
        f"/api/v1/assignments/{assignment['id']}/files",
        headers=student_headers,
        files={"file": ("report.md", f"# Report\n\n![pixel](data:image/png;base64,{png})".encode(), "text/markdown")},
    )
    assert markdown.status_code == 201, markdown.text
    files.append(markdown.json())
    assert student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "portfolio-current"}, json={}).status_code == 201
    teacher.post(f"/api/v1/assignments/{assignment['id']}/submissions/{member['id']}/grade", headers=headers, json={"grade": "C", "comment": "<p><strong>完整</strong><br>继续保持</p>"})
    with SessionLocal() as db:
        db.get(Assignment, UUID(other["id"])).due_at = datetime.now(UTC) - timedelta(days=1)
        version = db.scalar(select(SubmissionVersion).join(Submission).where(Submission.assignment_id == UUID(assignment['id'])))
        db.add(SubmissionAssessment(assignment_id=UUID(assignment["id"]), submission_version_id=version.id, evaluator_id=UUID(removed["id"]), subject_user_id=UUID(member["id"]), kind="PEER", grade="A", comment="<p>结构<strong>清晰</strong></p>", status="PUBLISHED"))
        db.commit()
    base = f"/api/v1/classes/{cid}/members/{member['id']}/portfolio"
    data = teacher.get(base).json()
    assert data["summary"] == {"total": 2, "submitted": 1, "late": 0, "missing": 1}
    submitted = next(item for item in data["assignments"] if item["status"] == "SUBMITTED")
    assert submitted["final_grade"] == "C" and submitted["peer_grade"] == "A"
    assert len(submitted["received_reviews"]) == 1
    assert next(item for item in data["assignments"] if item["status"] != "SUBMITTED")["grade_source"] == "SYSTEM"
    assert student.get(base).status_code == 403
    assert student.post(base + ".zip", headers=student_headers).status_code == 403
    assert TestClient(app).get(base).status_code == 401
    another = teacher.post("/api/v1/classes", headers=headers, json={"semester": "2039 秋季", "name": "另一班"}).json()
    assert teacher.get(f"/api/v1/classes/{another['id']}/members/{member['id']}/portfolio").status_code == 404
    def completed_archive(response):
        assert response.status_code == 202, response.text
        with SessionLocal() as db:
            job = db.get(BackgroundJob, UUID(response.json()["id"]))
            payload = job.payload
        process_archive_export(job.id, payload)
        key = next(key for key in isolated_object_storage if key.startswith(f"exports/{job.id}/"))
        return isolated_object_storage[key]

    response = teacher.post(base + ".zip", headers=headers)
    assert response.json()["filename"] == "20390001_=SUM(1,2).zip"
    with ZipFile(BytesIO(completed_archive(response))) as archive:
        assert len(set(archive.namelist())) == len(archive.namelist())
        markdown_name = next(name for name in archive.namelist() if name.endswith("report.md"))
        exported_markdown = archive.read(markdown_name).decode("utf-8")
        assert "images/" in exported_markdown and "/api/v1/markdown-assets/" not in exported_markdown
        assert any(name.endswith(".png") and "/images/" in name for name in archive.namelist())
        workbook = load_workbook(BytesIO(archive.read("学生档案.xlsx")))
        assert workbook.sheetnames == ["基本信息", "作业记录", "收到的互评"]
        assert workbook["基本信息"]["B2"].data_type == "s"
        assert workbook["作业记录"].max_row == 3
        assignment_headers = [cell.value for cell in workbook["作业记录"][1]]
        assert assignment_headers == ["作业", "截止时间", "提交状态", "提交时间", "提交版本", "迟交", "附件", "互评成绩", "教师评分", "最终成绩", "成绩来源", "评分状态", "教师评语"]
        assignment_row = next(row for row in workbook["作业记录"].iter_rows(min_row=2) if row[0].value == "个人报告")
        assert assignment_row[7].value == "A"
        assert assignment_row[8].value == "C"
        assert assignment_row[9].value == "C"
        assert assignment_row[12].value == "完整\n继续保持"
        assert workbook["收到的互评"]["D2"].value == "结构清晰"
        assert all("<" not in str(cell.value or "") for sheet in workbook for row in sheet for cell in row)
    assert teacher.delete(f"/api/v1/classes/{cid}/members/{removed['id']}", headers=headers).status_code == 204
    response = teacher.post(f"/api/v1/classes/{cid}/portfolio.zip", headers=headers)
    with ZipFile(BytesIO(completed_archive(response))) as archive:
        assert all(name.startswith("20390001_=SUM(1,2)/") for name in archive.namelist())
        assert not any(str(member["id"]) in name for name in archive.namelist())
        assert sum(name.endswith("/report.pdf") or "/report (" in name for name in archive.namelist()) == 2
        assert any(name.endswith("/学生档案.xlsx") for name in archive.namelist())
        assert any(name.endswith(".png") and "/images/" in name for name in archive.namelist())
    for file in files:
        assert student.delete(f"/api/v1/files/{file['id']}", headers=student_headers).status_code == 204
    replacement = student.post(f"/api/v1/assignments/{assignment['id']}/files", headers=student_headers, files={"file": ("updated.pdf", b"updated", "application/pdf")}).json()
    assert student.post(f"/api/v1/assignments/{assignment['id']}/submission", headers={**student_headers, "Idempotency-Key": "portfolio-v2"}, json={}).status_code == 201
    updated = teacher.get(base).json()
    current = next(item for item in updated["assignments"] if item["status"] == "SUBMITTED")
    assert current["version_no"] == 2 and current["final_grade"] == "C"
    assert current["grade_carried_forward"] is True and current["grading_status"] == "PENDING_REASSESSMENT"
    assert current["received_reviews"][0]["grade"] == "A"
    with ZipFile(BytesIO(completed_archive(teacher.post(base + ".zip", headers=headers)))) as archive:
        assert len(archive.namelist()) == 2
        assert any(name.endswith("/updated.pdf") for name in archive.namelist())
    with SessionLocal() as db:
        db.get(FileObject, UUID(replacement["id"])).storage_path = "missing-portfolio.pdf"
        db.commit()
    failed = teacher.post(base + ".zip", headers=headers)
    with SessionLocal() as db:
        job = db.get(BackgroundJob, UUID(failed.json()["id"]))
        payload = job.payload
    process_archive_export(job.id, payload)
    with SessionLocal() as db:
        assert db.get(BackgroundJob, job.id).status == "FAILED"


def test_portfolio_filename_safety():
    assert "/" not in safe_name("../../report.pdf")
    assert "\\" not in safe_name("..\\report.pdf")
    assert safe_name("CON.pdf") == "_CON.pdf"
    assert safe_name("...") == "unnamed"
