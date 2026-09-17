import io
import re
import zipfile
from datetime import datetime
from html import unescape
from pathlib import Path
from tempfile import TemporaryFile
from urllib.parse import quote
from uuid import UUID

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from app.models import Assignment, ClassMember, FileObject, Grade, PeerReview, ReviewCampaign, SubmissionAssessment, User, VersionFile
from app.settings import settings


def plain_text(value):
    separated = re.sub(r"</?(?:p|div|li|br|h[1-6]|blockquote|tr)[^>]*>", "\n", value or "", flags=re.IGNORECASE)
    from bleach import clean
    return re.sub(r"[ \t]+", " ", unescape(clean(separated, tags=[], strip=True))).strip()


def portfolio(db, cid, uid):
    from app.main import clean_html, file_json, latest_personal_submission, member_detail, missing_submission_grade_result, submission_grade_result
    person = member_detail(db, cid, uid)[2]
    assignments = db.scalars(select(Assignment).where(Assignment.class_id == cid, Assignment.submitter_type == "INDIVIDUAL", Assignment.status.in_(["PUBLISHED", "CLOSED"])).order_by(Assignment.due_at, Assignment.id)).all()
    items = []
    for assignment in assignments:
        submitted = latest_personal_submission(db, assignment.id, uid)
        version = submitted[1] if submitted else None
        result = submission_grade_result(db, version) if version else missing_submission_grade_result(assignment)
        files = db.scalars(select(FileObject).join(VersionFile, VersionFile.file_id == FileObject.id).where(VersionFile.version_id == version.id)).all() if version else []
        reviews = list(result["peer_feedbacks"])
        for review in reviews:
            review["comment"] = clean_html(review["comment"])
        if result["teacher_grade"]:
            result["teacher_grade"]["comment"] = clean_html(result["teacher_grade"]["comment"])
        legacy = db.scalar(select(Grade).where(Grade.assignment_id == assignment.id, Grade.subject_user_id == uid))
        legacy_result = None
        has_assessment = version and db.scalar(select(SubmissionAssessment.id).where(SubmissionAssessment.submission_version_id == version.id).limit(1))
        if legacy and not has_assessment:
            legacy_result = {"peer_score": float(legacy.peer_score) if legacy.peer_score is not None else None, "score": float(legacy.score) if legacy.score is not None else None, "draft_score": float(legacy.draft_score) if legacy.draft_score is not None else None, "status": legacy.status}
            if version:
                result = {**result, "final_grade": None, "grade_source": None}
        for review, reviewer in db.execute(select(PeerReview, User).join(User, User.id == PeerReview.reviewer_id).join(ReviewCampaign, ReviewCampaign.id == PeerReview.campaign_id).where(ReviewCampaign.assignment_id == assignment.id, PeerReview.reviewee_id == uid, PeerReview.status == "VALID")):
            reviews.append({"id": str(review.id), "evaluator_name": reviewer.display_name, "score": review.total_score, "comment": review.comment, "updated_at": review.updated_at, "submission_version_id": str(review.submission_version_id)})
        items.append({"assignment_id": str(assignment.id), "submission_version_id": str(version.id) if version else None, "title": assignment.title, "due_at": assignment.due_at, "status": "SUBMITTED" if version else "NOT_SUBMITTED", "submitted_at": version.submitted_at if version else None, "version_no": version.version_no if version else None, "is_late": version.is_late if version else False, "files": [file_json(file) for file in files], "received_reviews": reviews, "legacy_grade": legacy_result, **result})
    submitted_count = sum(item["status"] == "SUBMITTED" for item in items)
    return {"member": person, "assignments": items, "summary": {"total": len(items), "submitted": submitted_count, "late": sum(item["is_late"] for item in items), "missing": len(items) - submitted_count}}


def safe_name(value):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")[:100] or "unnamed"
    if name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        name = "_" + name
    return name


def workbook_bytes(data):
    from app.main import export_cell
    workbook = Workbook()
    basic = workbook.active
    basic.title = "基本信息"
    member = data["member"]
    basic.append(["学号", "姓名", "小组", "入班时间"])
    basic.append([member["student_no"], member["name"], member["team"] or "未入组", export_cell(datetime.fromisoformat(member["joined_at"]))])
    sheet = workbook.create_sheet("作业记录")
    sheet.append(["作业", "截止时间", "提交状态", "提交时间", "提交版本", "迟交", "附件", "互评成绩", "教师评分", "最终成绩", "成绩来源", "评分状态", "教师评语"])
    reviews = workbook.create_sheet("收到的互评")
    reviews.append(["作业", "评价人", "评价结果", "评语", "评价时间"])
    for item in data["assignments"]:
        final_score = item["final_grade"]
        source = {"TEACHER": "教师评分", "PEER": "学生互评", "SYSTEM": "系统判定"}.get(item["grade_source"], "")
        status = "已评分" if final_score is not None else "待评分" if item["status"] == "SUBMITTED" else "待提交"
        teacher_comment = plain_text((item["teacher_grade"] or {}).get("comment", ""))
        teacher_grade = (item["teacher_grade"] or {}).get("grade")
        sheet.append([item["title"], export_cell(item["due_at"]), "已提交" if item["status"] == "SUBMITTED" else "未提交", export_cell(item["submitted_at"]), item["version_no"], "是" if item["is_late"] else "否", "、".join(file["name"] for file in item["files"]), item["peer_grade"], teacher_grade, final_score, source, status, teacher_comment])
        for review in item["received_reviews"]:
            result = review.get("grade") if review.get("grade") is not None else review.get("score")
            reviews.append([item["title"], review["evaluator_name"], result, plain_text(review["comment"]), export_cell(review.get("published_at") or review["updated_at"])])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(50, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def write_student_archive(bundle, db, data, root=""):
    from app.main import ApiError
    bundle.writestr(root + "学生档案.xlsx", workbook_bytes(data))
    for index, item in enumerate(data["assignments"], 1):
        directory = root + f"{index:03d}_{safe_name(item['title'])}/"
        used = set()
        for file_info in item["files"]:
            file = db.get(FileObject, UUID(file_info["id"]))
            path = (settings.file_root / file.storage_path).resolve()
            if not path.is_relative_to(settings.file_root.resolve()) or not path.is_file():
                raise ApiError(404, "FILE_MISSING", "附件存储不可用，导出已取消")
            original = safe_name(file.original_name)
            name, suffix = original, 1
            while name.casefold() in used:
                name = f"{Path(original).stem} ({suffix}){Path(original).suffix}"
                suffix += 1
            used.add(name.casefold())
            bundle.write(path, directory + name)


def export_portfolios(db, course, user, uid=None):
    from app.main import ApiError, audit
    ids = [uid] if uid else db.scalars(select(ClassMember.user_id).where(ClassMember.class_id == course.id, ClassMember.status == "ACTIVE", ClassMember.role == "STUDENT").order_by(ClassMember.user_id)).all()
    archive = TemporaryFile()
    exported_member = None
    try:
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as bundle:
            for student_id in ids:
                data = portfolio(db, course.id, student_id)
                member = data["member"]
                exported_member = member if uid is not None else exported_member
                root = f"{safe_name(member['student_no'])}_{safe_name(member['name'])}/" if uid is None else ""
                write_student_archive(bundle, db, data, root)
        archive.seek(0)
        audit(db, user, "STUDENT_PORTFOLIOS_EXPORTED", "class", str(course.id), {"student_id": str(uid) if uid else None, "student_count": len(ids)})
        db.commit()
    except Exception:
        archive.close()
        raise

    def chunks():
        try:
            while chunk := archive.read(1024 * 1024):
                yield chunk
        finally:
            archive.close()

    filename = f"{safe_name(exported_member['student_no'])}_{safe_name(exported_member['name'])}.zip" if exported_member else f"{safe_name(course.name)}-班级档案.zip"
    return StreamingResponse(chunks(), media_type="application/zip", headers={"Content-Disposition": f"attachment; filename=portfolio.zip; filename*=UTF-8''{quote(filename)}"})
