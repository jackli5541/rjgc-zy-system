from __future__ import annotations

import csv
import io
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from uuid import UUID

from app.models import ClassMember, TeachingClass, Team, TeamMember, TeamRequest, Topic, User
from app.core.audit import audit
from app.core.deps import membership, require_team_window, require_writable_class
from app.core.errors import ApiError
from app.core.utils import now

def class_json(x: TeachingClass, *, member_count: int | None = None, assignment_count: int | None = None, deletable: bool | None = None):
    item = {"id": str(x.id), "course": x.course, "semester": x.semester, "name": x.name, "invite_code": x.invite_code, "status": x.status, "team_deadline": x.team_deadline, "topic_public": x.topic_public, "invite_requires_approval": x.invite_requires_approval, "version": x.version}
    if member_count is not None: item["member_count"] = member_count
    if assignment_count is not None: item["assignment_count"] = assignment_count
    if deletable is not None: item["deletable"] = deletable
    return item


def team_payload(x: Team, viewer: User, leader_name: str, member_count: int, topic: Topic | None, pending_count: int, topic_public: bool, own_team_id: UUID | None):
    can_view_topic = viewer.role == "TEACHER" or own_team_id == x.id or topic_public
    return {"id": str(x.id), "name": x.name, "leader_id": str(x.leader_id), "leader_name": leader_name, "member_count": member_count, "open_recruitment": x.open_recruitment, "status": x.status, "is_leader": x.leader_id == viewer.id, "pending_count": pending_count, "topic": None if not topic or not can_view_topic else {"id": str(topic.id), "name": topic.name, "description": topic.description, "status": topic.review_status, "reason": topic.review_reason}, "version": x.version}


def team_json(db: Session, x: Team, viewer: User):
    leader = db.get(User, x.leader_id)
    count = db.scalar(select(func.count()).select_from(TeamMember).where(TeamMember.team_id == x.id, TeamMember.status == "ACTIVE")) or 0
    topic = db.scalar(select(Topic).where(Topic.team_id == x.id))
    pending = db.scalar(select(func.count()).select_from(TeamRequest).where(TeamRequest.team_id == x.id, TeamRequest.status == "PENDING")) or 0
    course = db.get(TeachingClass, x.class_id)
    own_team = membership(db, x.class_id, viewer.id) if viewer.role == "STUDENT" else None
    return team_payload(x, viewer, leader.display_name, count, topic, pending, bool(course and course.topic_public), own_team[1].id if own_team else None)


def parse_roster(content: bytes, filename: str):
    if filename.lower().endswith(".csv"):
        return [(str(r.get("学号", "")).strip(), str(r.get("姓名", "")).strip()) for r in csv.DictReader(io.StringIO(content.decode("utf-8-sig")))]
    if filename.lower().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()
            row_iter = sheet.iter_rows(values_only=True)
            first_row = next(row_iter, None)
            headers = [str(x or "").strip() for x in first_row or ()]
            if "学号" not in headers or "姓名" not in headers: continue
            sid, name = headers.index("学号"), headers.index("姓名")
            return [(str(row[sid] or "").strip(), str(row[name] or "").strip()) for row in row_iter if len(row) > max(sid, name) and any(value is not None and str(value).strip() for value in row)]
        raise ValueError("missing roster headers")
    raise ApiError(422, "ROSTER_FORMAT_INVALID", "仅支持 XLSX 或 CSV 名单")


def member_detail(db: Session, cid: UUID, uid: UUID):
    row = db.execute(select(ClassMember, User).join(User).where(ClassMember.class_id == cid, ClassMember.user_id == uid, ClassMember.status == "ACTIVE")).first()
    if not row: raise ApiError(404, "MEMBER_NOT_FOUND", "未找到学生")
    member, student = row; team = membership(db, cid, uid)
    return member, student, {"id": str(student.id), "student_no": student.login_name, "name": student.display_name, "status": member.status, "team": team[1].name if team else None, "joined_at": member.joined_at.isoformat()}


def set_team_recruitment(tid: UUID, open_recruitment: bool, user: User, db: Session):
    team = db.scalar(select(Team).where(Team.id == tid).with_for_update())
    if not team or team.status != "ACTIVE": raise ApiError(404, "TEAM_NOT_FOUND", "小组不存在")
    if user.role != "STUDENT" or team.leader_id != user.id:
        raise ApiError(403, "TEAM_LEADER_REQUIRED", "仅组长可修改招募状态")
    course = require_writable_class(db, user, team.class_id)
    if open_recruitment: require_team_window(course, user)
    if not open_recruitment:
        db.execute(TeamRequest.__table__.update().where(TeamRequest.team_id == tid, TeamRequest.status == "PENDING").values(status="CANCELLED", resolved_at=now()))
    if team.open_recruitment != open_recruitment:
        team.open_recruitment = open_recruitment
        team.version += 1
        audit(db, user, "TEAM_RECRUITMENT_OPENED" if open_recruitment else "TEAM_RECRUITMENT_CLOSED", "team", str(tid))
    db.commit()
    return team_json(db, team, user)
